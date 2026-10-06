"""P27: ready-made Excel forms any company can add in one click: an expense claim, a travel
claim, a petty cash record, a leave application, an item request, a monthly attendance
record and a yearly stock count. Built here (openpyxl) in the company's language, with its
name on top, totals as formulas and space to sign; people download, fill and hand them back,
or ask their AI to fill them.

Each builder writes into a fresh workbook through `Sheet`, which keeps the look consistent:
a title, a company line, labelled header fields ("Name :" then a blank cell), one bordered
table with a shaded header and a total row, and signature boxes.
"""

import io
from collections.abc import Callable
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

THIN = Side(style="thin", color="8A8A8A")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINE = Border(bottom=THIN)
HEAD = PatternFill("solid", fgColor="E8F3EE")
BOLD = Font(bold=True)


class Sheet:
    def __init__(self, title: str, company: str, widths: list[int]):
        self.wb = Workbook()
        ws = self.wb.active
        assert isinstance(ws, Worksheet)
        self.ws: Worksheet = ws
        self.ws.title = title[:31]
        for i, w in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.cols = len(widths)
        self.row = 1
        self._merge_text(company.upper(), Font(bold=True, size=12))
        self._merge_text(title.upper(), Font(bold=True, size=13))
        self.row += 1

    def _merge_text(self, text: str, font: Font) -> None:
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.cols)
        c = ws.cell(row=self.row, column=1, value=text)
        c.font = font
        c.alignment = Alignment(horizontal="center")
        self.row += 1

    def fields(self, labels: list[str]) -> None:
        """Two labelled blanks per line: "Label :" then an underlined cell."""
        ws = self.ws
        half = max(self.cols // 2, 2)
        for i in range(0, len(labels), 2):
            for j, label in enumerate(labels[i : i + 2]):
                start = 1 + j * half
                ws.cell(row=self.row, column=start, value=label).font = BOLD
                ws.cell(row=self.row, column=start + 1, value=":")
                end = min(start + half - 1, self.cols)
                if end > start + 2:
                    ws.merge_cells(
                        start_row=self.row, start_column=start + 2, end_row=self.row, end_column=end
                    )
                for col in range(start + 2, end + 1):
                    ws.cell(row=self.row, column=col).border = LINE
            self.row += 1
        self.row += 1

    def table(
        self,
        heads: list[str],
        rows: int,
        *,
        numbered: bool = True,
        total: dict[int, str] | None = None,
        formula: Callable[[int], dict[int, str]] | None = None,
    ) -> int:
        """A bordered table; `total` = {column: label for the total row}; `formula(row)` gives
        per-row formulas by column. Returns the header row."""
        ws = self.ws
        head = self.row
        for i, h in enumerate(heads, start=1):
            c = ws.cell(row=head, column=i, value=h)
            c.font = BOLD
            c.fill = HEAD
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[head].height = 30
        for r in range(head + 1, head + 1 + rows):
            for i in range(1, len(heads) + 1):
                ws.cell(row=r, column=i).border = BOX
            if numbered:
                ws.cell(row=r, column=1, value=r - head).alignment = Alignment(horizontal="center")
            for col, f in (formula(r) if formula else {}).items():
                ws.cell(row=r, column=col, value=f)
        self.row = head + 1 + rows
        if total:
            first = min(total)
            ws.cell(row=self.row, column=max(first - 1, 1), value=total[first]).font = BOLD
            for col in total:
                letter = get_column_letter(col)
                c = ws.cell(
                    row=self.row,
                    column=col,
                    value=f"=SUM({letter}{head + 1}:{letter}{head + rows})",
                )
                c.font = BOLD
                c.border = BOX
                c.number_format = "#,##0.00"
            self.row += 1
        self.row += 1
        return head

    def money(self, col: int, first: int, last: int) -> None:
        for r in range(first, last + 1):
            self.ws.cell(row=r, column=col).number_format = "#,##0.00"

    def note(self, text: str) -> None:
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.cols)
        c = ws.cell(row=self.row, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(italic=True, size=9)
        ws.row_dimensions[self.row].height = 30
        self.row += 2

    def signatures(self, roles: list[str], name: str, date_label: str) -> None:
        ws = self.ws
        span = max(self.cols // max(len(roles), 1), 1)
        for i, role in enumerate(roles):
            col = 1 + i * span
            ws.cell(row=self.row, column=col, value=role).font = BOLD
            ws.cell(row=self.row + 3, column=col).border = LINE
            ws.cell(row=self.row + 4, column=col, value=f"{name} :")
            ws.cell(row=self.row + 5, column=col, value=f"{date_label} :")
        self.row += 7

    def bytes(self) -> bytes:
        self.ws.page_setup.orientation = "landscape" if self.cols > 7 else "portrait"
        self.ws.page_setup.fitToWidth = 1
        self.ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        out = io.BytesIO()
        self.wb.save(out)
        return out.getvalue()


# ---------------------------------------------------------------- the words, in two languages

W: dict[str, dict[str, str]] = {
    "en": {
        "name": "Name",
        "staff_no": "Staff no.",
        "dept": "Department",
        "month": "Month",
        "date": "Date",
        "no": "No.",
        "desc": "Description",
        "receipt": "Receipt no.",
        "amount": "Amount (RM)",
        "total": "TOTAL",
        "claimant": "Claimed by",
        "checked": "Checked by",
        "approved": "Approved by",
        "sig_name": "Name",
        "sig_date": "Date",
        "from": "From",
        "to": "To",
        "purpose": "Purpose",
        "km": "Km",
        "rate": "Rate per km (RM)",
        "toll": "Toll / parking (RM)",
        "opening": "Opening balance (RM)",
        "item": "Item",
        "cash_in": "Cash in (RM)",
        "cash_out": "Cash out (RM)",
        "balance": "Balance (RM)",
        "remarks": "Remarks",
        "leave_type": "Leave type",
        "days": "Days",
        "reason": "Reason",
        "start": "First day",
        "end": "Last day",
        "size": "Size",
        "qty": "Quantity",
        "site": "Site / location",
        "employee": "Employee",
        "worked": "Days worked",
        "ot": "Overtime (hours)",
        "absent": "Absent",
        "mc": "MC / leave",
        "location": "Location",
        "in_stock": "In stock",
        "good": "Good",
        "damaged": "Damaged",
        "need": "Needed",
        "receipts_note": "Attach every receipt. Claims without receipts are not paid.",
        "leave_note": "Hand in at least 3 working days before (emergency and MC: with the letter).",
        "phone": "Phone",
        "company": "Company",
    },
    "ms": {
        "name": "Nama",
        "staff_no": "No. pekerja",
        "dept": "Jabatan",
        "month": "Bulan",
        "date": "Tarikh",
        "no": "Bil.",
        "desc": "Perkara",
        "receipt": "No. resit",
        "amount": "Jumlah (RM)",
        "total": "JUMLAH",
        "claimant": "Dituntut oleh",
        "checked": "Disemak oleh",
        "approved": "Diluluskan oleh",
        "sig_name": "Nama",
        "sig_date": "Tarikh",
        "from": "Dari",
        "to": "Ke",
        "purpose": "Tujuan",
        "km": "Km",
        "rate": "Kadar sekm (RM)",
        "toll": "Tol / parkir (RM)",
        "opening": "Baki awal (RM)",
        "item": "Barang",
        "cash_in": "Wang masuk (RM)",
        "cash_out": "Wang keluar (RM)",
        "balance": "Baki (RM)",
        "remarks": "Catatan",
        "leave_type": "Jenis cuti",
        "days": "Hari",
        "reason": "Sebab",
        "start": "Hari pertama",
        "end": "Hari terakhir",
        "size": "Saiz",
        "qty": "Kuantiti",
        "site": "Lokasi / tapak",
        "employee": "Pekerja",
        "worked": "Hari bekerja",
        "ot": "Lebih masa (jam)",
        "absent": "Tidak hadir",
        "mc": "MC / cuti",
        "location": "Lokasi",
        "in_stock": "Ada dalam stok",
        "good": "Elok",
        "damaged": "Rosak",
        "need": "Perlu",
        "receipts_note": "Sertakan setiap resit. Tuntutan tanpa resit tidak akan dibayar.",
        "leave_note": (
            "Hantar sekurang-kurangnya 3 hari bekerja lebih awal (kecemasan dan MC: bersama surat)."
        ),
        "phone": "Telefon",
        "company": "Syarikat",
    },
}


@dataclass(frozen=True)
class Starter:
    key: str
    kind: str
    names: dict[str, str]
    descriptions: dict[str, str]
    schedule: dict[str, Any]
    build: Callable[[str, str], bytes]


def _expense(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(STARTERS_BY_KEY["expense_claim"].names[lang], company, [6, 14, 40, 16, 16, 22])
    s.fields([w["name"], w["staff_no"], w["dept"], w["month"]])
    head = s.table(
        [w["no"], w["date"], w["desc"], w["receipt"], w["amount"], w["remarks"]],
        15,
        total={5: w["total"]},
    )
    s.money(5, head + 1, head + 16)
    s.note(w["receipts_note"])
    s.signatures([w["claimant"], w["checked"], w["approved"]], w["sig_name"], w["sig_date"])
    return s.bytes()


def _travel(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(
        STARTERS_BY_KEY["travel_claim"].names[lang], company, [6, 12, 18, 18, 26, 8, 12, 14, 14]
    )
    s.fields([w["name"], w["staff_no"], w["dept"], w["month"]])

    def per_row(r: int) -> dict[int, str]:
        return {8: f'=IF(F{r}="","",F{r}*G{r}+IF(I{r}="",0,I{r}))'}

    head = s.table(
        [
            w["no"],
            w["date"],
            w["from"],
            w["to"],
            w["purpose"],
            w["km"],
            w["rate"],
            w["amount"],
            w["toll"],
        ],
        15,
        total={8: w["total"]},
        formula=per_row,
    )
    s.money(8, head + 1, head + 16)
    s.note(w["receipts_note"])
    s.signatures([w["claimant"], w["checked"], w["approved"]], w["sig_name"], w["sig_date"])
    return s.bytes()


def _petty(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(STARTERS_BY_KEY["petty_cash"].names[lang], company, [6, 12, 36, 14, 14, 14, 22])
    s.fields([w["month"], w["dept"]])

    def per_row(r: int) -> dict[int, str]:
        # Running balance: the row above plus money in, minus money out.
        return {6: f'=IF(AND(D{r}="",E{r}=""),"",N(F{r - 1})+N(D{r})-N(E{r}))'}

    heads = [w["no"], w["date"], w["item"], w["cash_in"], w["cash_out"], w["balance"], w["remarks"]]
    head = s.table(heads, 20, total={4: w["total"], 5: w["total"]}, formula=per_row)
    # The first row is the opening balance (money in), so the balance starts from it.
    first = head + 1
    s.ws.cell(row=first, column=3, value=w["opening"])
    s.ws.cell(
        row=first, column=6, value=f'=IF(AND(D{first}="",E{first}=""),"",N(D{first})-N(E{first}))'
    )
    for col in (4, 5, 6):
        s.money(col, head + 1, head + 21)
    s.note(w["receipts_note"])
    s.signatures([w["claimant"], w["checked"], w["approved"]], w["sig_name"], w["sig_date"])
    return s.bytes()


def _leave(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(STARTERS_BY_KEY["leave"].names[lang], company, [6, 22, 16, 16, 10, 40])
    s.fields([w["name"], w["staff_no"], w["dept"], w["phone"]])
    s.table(
        [w["no"], w["leave_type"], w["start"], w["end"], w["days"], w["reason"]],
        4,
        total={5: w["total"]},
    )
    s.note(w["leave_note"])
    s.signatures([w["claimant"], w["checked"], w["approved"]], w["sig_name"], w["sig_date"])
    return s.bytes()


def _request(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(STARTERS_BY_KEY["item_request"].names[lang], company, [6, 36, 12, 12, 36])
    s.fields([w["name"], w["staff_no"], w["site"], w["date"]])
    s.table([w["no"], w["item"], w["size"], w["qty"], w["reason"]], 12)
    s.signatures([w["claimant"], w["checked"], w["approved"]], w["sig_name"], w["sig_date"])
    return s.bytes()


def _attendance(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(STARTERS_BY_KEY["attendance"].names[lang], company, [6, 30, 14, 20, 14, 14, 14, 24])
    s.fields([w["month"], w["site"]])
    s.table(
        [
            w["no"],
            w["employee"],
            w["staff_no"],
            w["site"],
            w["worked"],
            w["ot"],
            w["mc"],
            w["remarks"],
        ],
        30,
        total={5: w["total"], 6: w["total"]},
    )
    s.signatures([w["claimant"], w["checked"]], w["sig_name"], w["sig_date"])
    return s.bytes()


def _stock(lang: str, company: str) -> bytes:
    w = W[lang]
    s = Sheet(STARTERS_BY_KEY["stock_count"].names[lang], company, [6, 30, 22, 12, 12, 12, 12, 24])
    s.fields([w["date"], w["site"]])
    s.table(
        [
            w["no"],
            w["item"],
            w["location"],
            w["in_stock"],
            w["good"],
            w["damaged"],
            w["need"],
            w["remarks"],
        ],
        30,
        total={4: w["total"]},
    )
    s.signatures([w["claimant"], w["checked"]], w["sig_name"], w["sig_date"])
    return s.bytes()


STARTERS: list[Starter] = [
    Starter(
        "expense_claim",
        "claim",
        {"en": "Expense claim", "ms": "Tuntutan perbelanjaan"},
        {
            "en": "Monthly claim for things bought for work, with receipts.",
            "ms": "Tuntutan bulanan bagi perbelanjaan kerja, bersama resit.",
        },
        {"every": "month", "from_day": 1, "to_day": 15},
        _expense,
    ),
    Starter(
        "travel_claim",
        "claim",
        {"en": "Travel claim", "ms": "Tuntutan perjalanan"},
        {
            "en": "Monthly mileage, toll and parking for work trips.",
            "ms": "Tuntutan bulanan jarak perjalanan, tol dan parkir untuk urusan kerja.",
        },
        {"every": "month", "from_day": 1, "to_day": 15},
        _travel,
    ),
    Starter(
        "petty_cash",
        "record",
        {"en": "Petty cash record", "ms": "Rekod petty cash"},
        {
            "en": "Money in and out of the office petty cash each month, with the balance.",
            "ms": "Wang masuk dan keluar petty cash pejabat setiap bulan, bersama baki.",
        },
        {"every": "month", "from_day": 1, "to_day": 5},
        _petty,
    ),
    Starter(
        "leave",
        "request",
        {"en": "Leave application", "ms": "Permohonan cuti"},
        {
            "en": "Apply for annual, emergency or medical leave.",
            "ms": "Mohon cuti tahunan, kecemasan atau sakit.",
        },
        {"every": "none"},
        _leave,
    ),
    Starter(
        "item_request",
        "request",
        {"en": "Item request", "ms": "Permohonan barang"},
        {
            "en": "Ask for uniforms, equipment or supplies, with sizes and quantities.",
            "ms": "Mohon uniform, peralatan atau bekalan, bersama saiz dan kuantiti.",
        },
        {"every": "none"},
        _request,
    ),
    Starter(
        "attendance",
        "report",
        {"en": "Monthly attendance record", "ms": "Rekod kehadiran bulanan"},
        {
            "en": "Days worked, overtime and leave for each person at a site.",
            "ms": "Hari bekerja, lebih masa dan cuti setiap pekerja di sesuatu lokasi.",
        },
        {"every": "month", "from_day": 28, "to_day": 3},
        _attendance,
    ),
    Starter(
        "stock_count",
        "checklist",
        {"en": "Yearly stock count", "ms": "Kiraan stok tahunan"},
        {
            "en": "Count every item at each site once a year: good, damaged, needed.",
            "ms": "Kira setiap barang di setiap lokasi sekali setahun: elok, rosak, perlu.",
        },
        {"every": "year", "month": 12, "from_day": 1, "to_day": 31},
        _stock,
    ),
]
STARTERS_BY_KEY = {s.key: s for s in STARTERS}


def build(key: str, lang: str, company: str) -> tuple[Starter, bytes]:
    s = STARTERS_BY_KEY[key]
    lang = lang if lang in W else "en"
    return s, s.build(lang, company)
