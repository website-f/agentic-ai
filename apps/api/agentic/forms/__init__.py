"""P27: company forms people fill in and hand back.

When a form is due (`window`): every month between two days (the 1st to the 15th; a window
that wraps, the 30th to the 3rd, runs into the next month), every year in one month, once
by a date, or whenever needed. A period names one round ("2026-10", "2026", "once"); each
person hands in one submission per period (a "whenever needed" form gets a new period each
time).

Filling an Excel form (`describe`, `fill`): the AI reads the layout as labelled cells
("B5 'NAMA PENGAWAL'", "D12 'KUANTITI'") and writes values into a copy, by cell, by label (the
first empty cell to the right of it) or as table rows under a header row. Formatting, merged
cells and formulas are kept; only values are written.

P29 limits: a workbook is looked inside before it is opened (documents.extract
.check_office_zip, FORM_MAX_UNZIPPED), only the first MAX_SCAN_ROWS x MAX_SCAN_COLS of a
sheet's used range are scanned for labels and headers, and text the AI writes that starts
like a formula (= + - @) is stored as text, never as a live formula.
"""

import io
import re
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, timedelta
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

KINDS = ("claim", "advance", "payroll", "request", "report", "record", "checklist", "other")
EXCEL = (".xlsx", ".xlsm")
FORM_MAX_UNZIPPED = 40 * 1024 * 1024  # a form, not a database: openpyxl holds every cell
MAX_SCAN_ROWS = 2000
MAX_SCAN_COLS = 200
FORMULA_START = ("=", "+", "-", "@")


# ---------------------------------------------------------------- when it is due


@dataclass(frozen=True)
class Window:
    period: str
    opens: date
    due: date


def _day(y: int, m: int, d: int) -> date:
    return date(y, m, min(max(int(d), 1), monthrange(y, m)[1]))


def _month(y: int, m: int, k: int) -> tuple[int, int]:
    n = y * 12 + (m - 1) + k
    return n // 12, n % 12 + 1


def window(schedule: dict[str, Any] | None, today: date) -> Window | None:
    """The round of a form that is open now or comes next (None: whenever needed)."""
    s = schedule or {}
    every = s.get("every") or "none"
    if every == "month":
        f, t = int(s.get("from_day") or 1), int(s.get("to_day") or 28)
        for k in (-1, 0, 1):
            y, m = _month(today.year, today.month, k)
            opens = _day(y, m, f)
            ey, em = _month(y, m, 1) if t < f else (y, m)
            due = _day(ey, em, t)
            if today <= due:
                return Window(f"{y:04d}-{m:02d}", opens, due)
    if every == "year":
        mo = min(max(int(s.get("month") or 12), 1), 12)
        f, t = int(s.get("from_day") or 1), int(s.get("to_day") or 31)
        for y in (today.year, today.year + 1):
            opens = _day(y, mo, f)
            ey, em = _month(y, mo, 1) if t < f else (y, mo)
            due = _day(ey, em, t)
            if today <= due:
                return Window(f"{y:04d}", opens, due)
    if every == "once":
        try:
            due = date.fromisoformat(str(s.get("due_on")))
        except ValueError:
            return None
        return Window("once", min(today, due), due)
    return None


LATE_DAYS = 10


def missed(schedule: dict[str, Any] | None, today: date) -> Window | None:
    """The round whose deadline passed in the last LATE_DAYS days (still worth handing in)."""
    s = schedule or {}
    if (s.get("every") or "none") not in ("month", "year", "once"):
        return None
    w = window(s, today - timedelta(days=LATE_DAYS))
    if w is not None and w.due < today <= w.due + timedelta(days=LATE_DAYS):
        return w
    return None


def state(w: Window | None, today: date, submitted: str | None) -> str:
    """For one person: accepted | submitted | returned | draft (theirs), else upcoming |
    open | due_soon | late, or anytime for a form with no schedule."""
    if submitted in ("accepted", "submitted", "returned", "draft"):
        return submitted
    if w is None:
        return "anytime"
    if today < w.opens:
        return "upcoming"
    if today > w.due:
        return "late"
    return "due_soon" if (w.due - today) <= timedelta(days=2) else "open"


def describe_schedule(schedule: dict[str, Any] | None) -> str:
    """Plain English for the AI's brief (the app shows its own, translated)."""
    s = schedule or {}
    every = s.get("every") or "none"
    if every == "month":
        return f"every month, from day {s.get('from_day', 1)} to day {s.get('to_day', 28)}"
    if every == "year":
        return (
            f"every year in month {s.get('month', 12)}, "
            f"day {s.get('from_day', 1)} to {s.get('to_day', 31)}"
        )
    if every == "once":
        return f"once, by {s.get('due_on')}"
    return "whenever needed"


def clean_schedule(raw: dict[str, Any] | None) -> dict[str, Any]:
    s = raw or {}
    every = s.get("every") if s.get("every") in ("month", "year", "once", "none") else "none"
    out: dict[str, Any] = {"every": every}
    if every in ("month", "year"):
        out["from_day"] = min(max(int(s.get("from_day") or 1), 1), 31)
        out["to_day"] = min(max(int(s.get("to_day") or 28), 1), 31)
    if every == "year":
        out["month"] = min(max(int(s.get("month") or 12), 1), 12)
    if every == "once":
        try:
            out["due_on"] = date.fromisoformat(str(s.get("due_on"))).isoformat()
        except ValueError:
            out = {"every": "none"}
    return out


# ---------------------------------------------------------------- reading and filling Excel


def is_excel(name: str) -> bool:
    return name.lower().endswith(EXCEL)


def _text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).split())


def _load(data: bytes) -> Any:
    """The workbook, after checking it does not unpack into more than a form's worth."""
    from ..documents.extract import check_office_zip

    check_office_zip(data, FORM_MAX_UNZIPPED)
    return load_workbook(io.BytesIO(data))


def _scan(ws: Worksheet) -> Any:
    """The rows of a sheet's used range, at most MAX_SCAN_ROWS x MAX_SCAN_COLS."""
    return ws.iter_rows(
        min_row=max(1, ws.min_row),
        max_row=min(ws.max_row, max(1, ws.min_row) + MAX_SCAN_ROWS - 1),
        min_col=max(1, ws.min_column),
        max_col=min(ws.max_column, max(1, ws.min_column) + MAX_SCAN_COLS - 1),
    )


def describe(data: bytes, limit: int = 450) -> str:
    """The form's layout for the AI: every sheet, its filled cells by reference (formulas
    shown as formulas), and the merged areas, so it can tell labels from blanks."""
    wb = _load(data)
    lines: list[str] = []
    n = 0
    for ws in wb.worksheets:
        lines.append(f"## Sheet {ws.title!r} ({ws.max_row} rows x {ws.max_column} columns)")
        for row in _scan(ws):
            for c in row:
                if isinstance(c, MergedCell) or c.value is None or _text(c.value) == "":
                    continue
                lines.append(f"{c.coordinate} {_text(c.value)[:80]!r}")
                n += 1
                if n >= limit:
                    lines.append("... (more cells not shown)")
                    return "\n".join(lines)
        merged = [str(r) for r in ws.merged_cells.ranges][:40]
        if merged:
            lines.append("Merged: " + ", ".join(merged))
    return "\n".join(lines)


def _sheet(wb: Any, name: str | None) -> Worksheet:
    if name and name in wb.sheetnames:
        return wb[name]
    return wb.worksheets[0]


def _writable(ws: Worksheet, ref: str) -> Any:
    """The cell to write for `ref`: a merged area writes into its top-left cell."""
    c = ws[ref]
    if isinstance(c, MergedCell):
        for r in ws.merged_cells.ranges:
            if ref in r:
                return ws.cell(row=r.min_row, column=r.min_col)
    return c


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _number(v: Any) -> Any:
    """Numbers stay numbers so the form's formulas add them up ("RM1,234.50" -> 1234.5)."""
    if isinstance(v, int | float):
        return v
    t = str(v).strip()
    m = re.fullmatch(r"(?:RM\s?)?(-?[\d,]+(?:\.\d+)?)", t, re.I)
    if m:
        n = float(m.group(1).replace(",", ""))
        return int(n) if n.is_integer() and "." not in m.group(1) else n
    return v


def _put(cell: Any, value: Any) -> None:
    """Write a value the AI gave. Text starting like a formula ("=HYPERLINK(...)", "+60...",
    "-5 days", "@x") is kept as text: written into the form it would run as a formula (or
    be read as one by whoever opens it). The form's own formulas are never overwritten
    (callers skip those cells)."""
    value = _number(value)
    cell.value = value
    if isinstance(value, str) and value.startswith(FORMULA_START):
        cell.data_type = "s"  # openpyxl: a string, saved as text even though it starts with =


def _find_label(ws: Worksheet, label: str) -> Any:
    want = _norm(label)
    best = None
    for row in _scan(ws):
        for c in row:
            if isinstance(c, MergedCell) or not isinstance(c.value, str):
                continue
            have = _norm(c.value)
            if have == want:
                return c
            if want and want in have and best is None:
                best = c
    return best


def _right_of(ws: Worksheet, c: Any) -> Any:
    """The first empty cell to the right of a label, skipping ":" cells (None when the next
    cell already holds something else)."""
    for col in range(c.column + 1, min(c.column + 12, ws.max_column + 2)):
        t = _writable(ws, f"{get_column_letter(col)}{c.row}")
        if t.coordinate == c.coordinate:  # still inside the label's merged area
            continue
        v = _text(t.value)
        if v in (":", "-"):
            continue
        return t if v == "" or set(v) <= {"_", "."} else None
    return None


def fill(
    data: bytes,
    *,
    cells: dict[str, Any] | None = None,
    fields: dict[str, Any] | None = None,
    rows: list[dict[str, Any]] | None = None,
    sheet: str | None = None,
) -> tuple[bytes, list[str]]:
    """A filled copy of the form and what was done (or could not be placed)."""
    wb = _load(data)
    ws = _sheet(wb, sheet)
    report: list[str] = []
    for ref, value in (cells or {}).items():
        ref = str(ref).strip().upper()
        target_ws = ws
        if "!" in ref:
            name, ref = ref.rsplit("!", 1)
            target_ws = _sheet(wb, name.strip("'"))
        try:
            coordinate_to_tuple(ref)
        except ValueError:
            report.append(f"skipped {ref!r}: not a cell reference")
            continue
        target = _writable(target_ws, ref)
        if target.data_type == "f" or (
            isinstance(target.value, str) and target.value.startswith("=")
        ):
            report.append(f"kept the formula in {ref}")
            continue
        _put(target, value)
    for label, value in (fields or {}).items():
        c = _find_label(ws, str(label))
        if c is None:
            report.append(f"no label {label!r} in the form")
            continue
        target = _right_of(ws, c)
        if target is None:
            report.append(f"no empty cell next to {label!r}")
            continue
        _put(target, value)
        report.append(f"{label} -> {target.coordinate}")
    if rows:
        report += _rows(ws, rows)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), report


def _rows(ws: Worksheet, rows: list[dict[str, Any]]) -> list[str]:
    """Write table rows under the header row that names most of their keys."""
    keys = {k for r in rows for k in r}
    best: tuple[int, int, dict[str, int]] | None = None
    for row in _scan(ws):
        cols: dict[str, int] = {}
        for c in row:
            if isinstance(c, MergedCell) or not isinstance(c.value, str):
                continue
            for k in keys:
                if k not in cols and _norm(k) and _norm(k) == _norm(c.value):
                    cols[k] = c.column
        if cols and (best is None or len(cols) > best[1]):
            best = (int(row[0].row or 0), len(cols), cols)
    if best is None:
        return [f"no header row with {', '.join(sorted(keys))}"]
    header, _, cols = best
    r = header + 1
    written = 0
    for item in rows:
        while r <= header + MAX_SCAN_ROWS and any(
            _text(_writable(ws, f"{get_column_letter(col)}{r}").value) for col in cols.values()
        ):
            r += 1
        for k, col in cols.items():
            if k in item:
                _put(_writable(ws, f"{get_column_letter(col)}{r}"), item[k])
        written += 1
        r += 1
    missing = sorted(keys - set(cols))
    out = [f"{written} rows under the header on row {header}"]
    if missing:
        out.append(f"no column for {', '.join(missing)}")
    return out
