"""Export a document's data to Excel: a Details sheet (number, date, company, fields) and one
sheet per table, with numbers stored as numbers so they sum and sort."""

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .blocks import parse, plain
from .render_pdf import hex_rgb

_NUM = re.compile(r"^-?(RM|USD|SGD|\$)?\s?(-?[\d,]+(\.\d+)?)$")


def _value(cell: str) -> Any:
    text = plain(cell).strip()
    m = _NUM.match(text)
    if m and len(m.group(2).replace(",", "").lstrip("-").split(".")[0]) <= 12:
        n = float(m.group(2).replace(",", ""))
        return int(n) if n.is_integer() and "." not in m.group(2) else n
    return text


def _sheet_name(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:28] or "Table"
    name, n = base, 2
    while name in used:
        name, n = f"{base[:25]} {n}", n + 1
    used.add(name)
    return name


def render(
    markdown: str,
    details: list[tuple[str, str]],
    title: str = "",
    accent: str = "",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Details"
    used = {"Details"}
    head_fill = PatternFill(
        "solid", fgColor="".join(f"{int(255 - (255 - v) * 0.15):02X}" for v in hex_rgb(accent))
    )
    ws["A1"] = title or "Document"
    ws["A1"].font = Font(bold=True, size=14)
    for i, (k, v) in enumerate(details, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70

    heading = ""
    for b in parse(markdown):
        if b.kind == "heading":
            heading = plain(b.lines[0])
            continue
        if b.kind != "table" or not b.header:
            continue
        sh = wb.create_sheet(_sheet_name(heading or "Table", used))
        for c, h in enumerate(b.header, 1):
            cell = sh.cell(1, c, plain(h))
            cell.font = Font(bold=True)
            cell.fill = head_fill
        widths = [len(plain(h)) for h in b.header]
        for r, row in enumerate(b.rows, 2):
            for c, v in enumerate(row, 1):
                val = _value(v)
                cell = sh.cell(r, c, val)
                if "**" in v:
                    cell.font = Font(bold=True)
                if isinstance(val, float):
                    cell.number_format = "#,##0.00"
                widths[c - 1] = max(widths[c - 1], len(str(val)))
        for c, w in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(c)].width = min(60, max(8, w + 2))
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
