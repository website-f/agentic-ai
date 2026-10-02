"""Render a document to Word (.docx) with python-docx: the same blocks and letterhead as the
PDF, as real Word structure (headings, lists, tables) so people can keep editing it."""

import io
from typing import Any

from docx import Document as Docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .blocks import SIGN_LINE, Block, is_numeric, parse, runs
from .render_pdf import Letterhead

INK = RGBColor(0x19, 0x19, 0x19)
GREY = RGBColor(0x5F, 0x5F, 0x5F)


def _hex(rgb: tuple[int, int, int]) -> str:
    return "".join(f"{v:02X}" for v in rgb)


def _shade(cell: Any, rgb: tuple[int, int, int]) -> None:
    tc = cell._tc.get_or_add_tcPr()  # noqa: SLF001 - python-docx has no shading API
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(rgb))
    tc.append(shd)


def _bottom_border(paragraph: Any, rgb: tuple[int, int, int], size: int = 8) -> None:
    ppr = paragraph._p.get_or_add_pPr()  # noqa: SLF001
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), _hex(rgb))
    bdr.append(bottom)
    ppr.append(bdr)


def _field(run: Any, code: str) -> None:
    """A Word field (PAGE / NUMPAGES) inside a run."""
    for kind, text in (("begin", None), (None, code), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)  # noqa: SLF001


def _add_runs(paragraph: Any, text: str, size: float = 10, italic: bool = False) -> None:
    for r in runs(text):
        run = paragraph.add_run(r.text)
        run.bold = r.bold
        run.italic = r.italic or italic
        run.font.size = Pt(size)


def _letterhead(doc: Any, lh: Letterhead) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    head = section.header
    p = head.paragraphs[0]
    if lh.logo:
        try:
            p.add_run().add_picture(io.BytesIO(lh.logo), height=Mm(14))
            p = head.add_paragraph()
        except Exception:  # noqa: BLE001, S110 - a bad logo never breaks the document
            pass
    name = p.add_run(lh.name)
    name.bold = True
    name.font.size = Pt(14)
    name.font.color.rgb = INK
    for line in lh.lines:
        q = head.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        r = q.add_run(line)
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
    _bottom_border(head.paragraphs[-1], lh.accent, 10)

    foot = section.footer.paragraphs[0]
    r = foot.add_run(lh.footer[:140] + ("    " if lh.footer else ""))
    r.font.size = Pt(7.5)
    r.font.color.rgb = GREY
    for label, code in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        t = foot.add_run(label)
        t.font.size = Pt(7.5)
        t.font.color.rgb = GREY
        f = foot.add_run()
        f.font.size = Pt(7.5)
        _field(f, code)


def _table(doc: Any, b: Block, accent: tuple[int, int, int]) -> None:
    cols = len(b.header)
    table = doc.add_table(rows=1 + len(b.rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    body = [r for r in b.rows if r and r[0].strip()] or b.rows
    right = []
    for c in range(cols):
        cells = [r[c] for r in body if c < len(r) and r[c].strip()]
        right.append(bool(cells) and sum(is_numeric(x) for x in cells) >= max(1, len(cells) * 0.6))
    light = tuple(int(255 - (255 - v) * 0.12) for v in accent)
    for ri, row in enumerate([b.header, *b.rows]):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            p = cell.paragraphs[0]
            _add_runs(p, row[ci] if ci < len(row) else "", 9)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                _shade(cell, light)  # type: ignore[arg-type]
            if right[ci]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def render(markdown: str, title: str = "", lh: Letterhead | None = None) -> bytes:
    doc = Docx()
    section = doc.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin, section.bottom_margin = Mm(16), Mm(18)
    normal: Any = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    doc.core_properties.title = title or "Document"
    if lh and lh.name:
        doc.core_properties.author = lh.name
        _letterhead(doc, lh)
    accent = lh.accent if lh else (19, 137, 95)

    for b in parse(markdown):
        if b.kind == "heading":
            h = doc.add_heading(level=b.level)
            _add_runs(h, b.lines[0], {1: 15, 2: 12.5, 3: 11}.get(b.level, 11))
            for run in h.runs:
                run.font.color.rgb = INK
                run.bold = True
        elif b.kind == "para":
            p = doc.add_paragraph()
            for i, line in enumerate(b.lines):
                if SIGN_LINE.match(line):
                    p.add_run().add_break()
                    p.add_run().add_break()
                elif i:
                    p.add_run().add_break()
                _add_runs(p, line)
        elif b.kind == "note":
            p = doc.add_paragraph()
            for i, line in enumerate(b.lines):
                if i:
                    p.add_run().add_break()
                _add_runs(p, line, 9.5, italic=True)
            for run in p.runs:
                run.font.color.rgb = GREY
        elif b.kind in ("bullets", "numbers"):
            style = "List Bullet" if b.kind == "bullets" else "List Number"
            for item in b.lines:
                _add_runs(doc.add_paragraph(style=style), item)
        elif b.kind == "rule":
            _bottom_border(doc.add_paragraph(), (200, 200, 200), 6)
        elif b.kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif b.kind == "table" and b.header:
            _table(doc, b, accent)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
