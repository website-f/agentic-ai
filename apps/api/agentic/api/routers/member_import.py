"""Add many members at once from an Excel or CSV sheet.

1. `GET  /api/members/import/template`: an .xlsx to fill in (and a sheet of the companies,
   departments and roles the reader may use).
2. `POST /api/members/import/preview`: the sheet as raw bytes (like /api/files; see
   RAW_UPLOAD_PATHS in api/main.py). Finds the header row wherever it is, maps the columns
   (English and Malay headers; the cheap AI model only when Name or Email cannot be found),
   matches each row's company and department to the ones that exist (never makes new ones)
   and checks each row with the same rules as adding one person. Nothing is saved.
3. `POST /api/members/import/check`: the rows again after the person fixed some (JSON).
4. `POST /api/members/import`: adds the rows that pass, through members.add_one, and hands
   back each new account's one-time password.
"""

import csv
import difflib
import io
import json
import logging
import re
from typing import Any, Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import Response
from pydantic import BaseModel, EmailStr, Field, TypeAdapter, ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ...core.db import get_db
from ...core.fence import fence
from ...core.security import ROLES, SCOPED_ROLES, can
from ...i18n import MS, current_lang, render, tr
from ...i18n import labels as i18n_labels
from ...models import Branch, Department, Membership, User
from ...services import audit
from ..deps import Principal, api_error
from . import members

router = APIRouter(prefix="/api/members/import", tags=["members"])
log = logging.getLogger("agentic.api.member_import")

UPLOAD_PATH = "/api/members/import/preview"
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
HEADER_SCAN = 20  # rows looked at for the header row
FIELDS = ("name", "email", "role", "company", "department", "phone", "notes")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Header words, English and Malay, normalised (see _norm). An exact match wins; otherwise
# the first field (in this order) whose word appears in the header: "Company email" is an
# email column, "Nama syarikat" a company column, and "name" is tried last.
HEADERS: dict[str, tuple[str, ...]] = {
    "email": ("email", "e mail", "emel", "e mel", "mail", "email address", "alamat emel"),
    "department": ("department", "dept", "jabatan", "bahagian", "unit", "division"),
    "company": ("company", "branch", "syarikat", "cawangan", "outlet", "organisation"),
    "role": ("role", "access", "position", "jawatan", "peranan", "akses", "level"),
    "phone": ("phone", "mobile", "tel", "telephone", "telefon", "no telefon", "h p", "hp"),
    "notes": ("notes", "note", "remarks", "nota", "catatan", "ulasan"),
    "name": ("name", "full name", "nama", "nama penuh", "staff name", "employee name"),
}

# Role words a sheet may use, normalised -> role key. The keys, the app's labels in both
# languages (added below) and what Malaysian offices call them.
ROLE_WORDS: dict[str, str] = {
    "pemilik": "owner",
    "tuan punya": "owner",
    "administrator": "admin",
    "pentadbir": "admin",
    "manager": "branch_manager",
    "pengurus": "branch_manager",
    "bm": "branch_manager",
    "hod": "hod",
    "head": "hod",
    "head of dept": "hod",
    "ketua": "hod",
    "ketua bahagian": "hod",
    "ketua unit": "hod",
    "penyelia": "supervisor",
    "staf": "staff",
    "pekerja": "staff",
    "employee": "staff",
    "worker": "staff",
    "user": "staff",
    "pengendali": "operator",
    "pelulus": "approver",
    "pemerhati": "viewer",
    "read only": "viewer",
    "view only": "viewer",
    "lihat sahaja": "viewer",
    "baca sahaja": "viewer",
}

# Malay department words -> the English names workspaces often use, to match either way.
DEPT_WORDS: dict[str, str] = {
    "kewangan": "finance",
    "akaun": "accounts",
    "perakaunan": "accounting",
    "operasi": "operations",
    "pengurusan": "management",
    "penyelidikan": "research",
    "penulisan": "writing",
    "sumber manusia": "human resources",
    "pemasaran": "marketing",
    "jualan": "sales",
    "pentadbiran": "administration",
    "keselamatan": "security",
    "perolehan": "procurement",
    "teknologi maklumat": "it",
    "khidmat pelanggan": "customer service",
    "undang undang": "legal",
}

# Company words that do not tell companies apart ("Maju Sdn Bhd" is "Maju").
COMPANY_NOISE = {"sdn", "bhd", "berhad", "sendirian", "plc", "ltd", "limited", "inc", "co", "m"}

_EMAIL = TypeAdapter(EmailStr)


def _norm(value: Any) -> str:
    return " ".join(re.sub(r"[^0-9a-z]+", " ", str(value or "").lower()).split())


def _role_words() -> dict[str, str]:
    words = {_norm(k): k for k in ROLES}
    for key, label in i18n_labels.ROLES.items():
        words[_norm(label)] = key
        words[_norm(MS.get(label, label))] = key
    words.update(ROLE_WORDS)
    return words


ROLE_LOOKUP = _role_words()


# ---------------------------------------------------------------- shapes


class Note(BaseModel):
    # What the note is about (name, email, role, company, department, place, member) and
    # where it came from: "match" (reading the sheet) or "check" (the rules for adding).
    field: str
    level: Literal["warn", "error"]
    source: Literal["match", "check"]
    text: str


class ImportRow(BaseModel):
    row: int = Field(ge=0)  # the sheet's row number, for people to find it
    name: str = Field(default="", max_length=200)
    email: str = Field(default="", max_length=320)
    role: str = Field(default="staff", max_length=40)
    branch_id: str | None = Field(default=None, max_length=40)
    department_id: str | None = Field(default=None, max_length=40)
    phone: str = Field(default="", max_length=60)
    notes: str = Field(default="", max_length=500)


class RowOut(ImportRow):
    branch_name: str | None = None
    department_name: str | None = None
    # What the sheet said, for rows the person fixes by hand.
    role_text: str = ""
    company_text: str = ""
    department_text: str = ""
    existing_account: bool = False
    status: Literal["ok", "warn", "error"] = "ok"
    notes_found: list[Note] = Field(default_factory=list)


class ColumnOut(BaseModel):
    field: str
    header: str
    column: int  # 0-based
    letter: str


class PreviewOut(BaseModel):
    file_name: str
    sheet: str
    header_row: int  # 1-based; 0 when the sheet has no header row the AI could find
    mapped_by: Literal["headers", "ai"]
    columns: list[ColumnOut]
    rows: list[RowOut]
    counts: dict[str, int]
    capped: bool  # more than MAX_ROWS people: only the first MAX_ROWS were read


class RowsIn(BaseModel):
    rows: list[ImportRow] = Field(min_length=1, max_length=MAX_ROWS)


class CheckOut(BaseModel):
    rows: list[RowOut]
    counts: dict[str, int]


class ResultOut(BaseModel):
    row: int
    name: str
    email: str
    status: Literal["added", "skipped"]
    message: str | None = None
    temp_password: str | None = None  # shown once; None for an existing account
    user_id: str | None = None
    role: str
    branch_name: str | None = None
    department_name: str | None = None


class ImportOut(BaseModel):
    results: list[ResultOut]
    added: int
    skipped: int


# ---------------------------------------------------------------- reading the sheet


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # a phone number Excel stored as a number
    return str(v).strip()


def _read_rows(data: bytes, name: str) -> tuple[str, list[list[str]]]:
    """The first sheet as rows of text (at most MAX_ROWS + HEADER_SCAN + 1 of them)."""
    limit = MAX_ROWS + HEADER_SCAN + 1
    if data[:4] == b"PK\x03\x04":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 - any broken zip or sheet reads the same
            raise api_error(
                status.HTTP_400_BAD_REQUEST,
                "bad_sheet",
                "That file could not be read as Excel. Save it as .xlsx or CSV and try again.",
            ) from e
        try:
            ws = wb.worksheets[0]
            rows: list[list[str]] = []
            for raw in ws.iter_rows(values_only=True):
                rows.append([_cell(v) for v in raw[:40]])
                if len(rows) >= limit:
                    break
            return ws.title, rows
        finally:
            wb.close()
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise api_error(
            status.HTTP_400_BAD_REQUEST,
            "old_excel",
            "This is an old Excel file (.xls). Save it as .xlsx or CSV and upload it again.",
        )
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any bytes
        text = ""
    if "\x00" in text[:2000]:
        raise api_error(
            status.HTTP_400_BAD_REQUEST,
            "bad_sheet",
            "That file could not be read as Excel. Save it as .xlsx or CSV and try again.",
        )
    sample = text[:4000]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for raw in csv.reader(io.StringIO(text), dialect):
        rows.append([c.strip() for c in raw[:40]])
        if len(rows) >= limit:
            break
    return name.rsplit(".", 1)[0][:60] or "CSV", rows


def header_field(header: str) -> str | None:
    """Which field a column header is about, or None."""
    h = _norm(header)
    if not h:
        return None
    for field, words in HEADERS.items():
        if h in words:
            return field
    padded = f" {h} "
    for field, words in HEADERS.items():
        if any(f" {w} " in padded for w in words):
            return field
    return None


def _map_header(row: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, cell in enumerate(row):
        field = header_field(cell)
        if field and field not in out:
            out[field] = i
    return out


def find_header(rows: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    """The header row (index) and its columns: the row among the first HEADER_SCAN that
    names the most fields, as long as it names the person's name or email."""
    best: tuple[int, dict[str, int]] | None = None
    for i, row in enumerate(rows[:HEADER_SCAN]):
        cols = _map_header(row)
        if ("name" in cols or "email" in cols) and (best is None or len(cols) > len(best[1])):
            best = (i, cols)
    return best


AI_COLUMNS = (
    "You map the columns of a spreadsheet that lists people to add to an office app. Answer "
    "JSON only: "
    '{"header_row": <0-based index of the header row, or -1 if there is none>, '
    '"columns": {"name": <0-based column or -1>, "email": ..., "role": ..., "company": ..., '
    '"department": ..., "phone": ..., "notes": ...}}. "company" is the company or branch the '
    "person works in; headers may be in English or Malay."
)


async def _ai_columns(
    db: AsyncSession, workspace_id: str, rows: list[list[str]]
) -> tuple[int, dict[str, int]] | None:
    """Ask the cheap model which column is which. Only the first rows go, cut short."""
    from ...engine import gateway

    sample = [[c[:40] for c in r[:15]] for r in rows[:8]]
    table = "\n".join(f"{i}: " + " | ".join(r) for i, r in enumerate(sample))
    try:
        r = await gateway.chat(
            db,
            workspace_id,
            "fast",
            [
                {"role": "system", "content": AI_COLUMNS},
                {"role": "user", "content": f"Rows:\n{fence(table)}"},
            ],
            task="members.import_columns",
            max_tokens=300,
            temperature=0,
            json_mode=True,
        )
        raw = r.content or ""
        start, end = raw.find("{"), raw.rfind("}")
        data = json.loads(raw[start : end + 1]) if start >= 0 else {}
    except Exception:  # noqa: BLE001 - no model, a bad reply: the person gets the plain error
        log.info("AI column mapping unavailable", exc_info=True)
        return None
    if not isinstance(data, dict):
        return None
    header = data.get("header_row")
    cols_in = data.get("columns")
    if not isinstance(header, int) or not isinstance(cols_in, dict):
        return None
    width = max((len(r) for r in rows[:HEADER_SCAN]), default=0)
    cols: dict[str, int] = {}
    for field in FIELDS:
        v = cols_in.get(field)
        if isinstance(v, int) and 0 <= v < width and v not in cols.values():
            cols[field] = v
    if "name" not in cols or "email" not in cols:
        return None
    return (header if 0 <= header < min(len(rows), HEADER_SCAN) else -1), cols


def _letter(i: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(i + 1)


# ---------------------------------------------------------------- matching


class Org:
    """The companies and departments this person may place people in, for matching."""

    def __init__(self, branches: list[Branch], departments: list[Department]):
        self.branches = {b.id: b for b in branches}
        self.departments = {d.id: d for d in departments}

    @classmethod
    async def load(cls, db: AsyncSession, workspace_id: str) -> "Org":
        branches = (
            await db.scalars(
                select(Branch).where(Branch.workspace_id == workspace_id).order_by(Branch.name)
            )
        ).all()
        depts = (
            await db.scalars(
                select(Department)
                .where(Department.workspace_id == workspace_id)
                .order_by(Department.name)
            )
        ).all()
        return cls(list(branches), list(depts))


def _company_key(name: str) -> str:
    words = _norm(name).split()
    return " ".join(w for w in words if w not in COMPANY_NOISE) or " ".join(words)


def _dept_key(name: str) -> str:
    key = _norm(name)
    return DEPT_WORDS.get(key, key)


def best_match(text: str, options: dict[str, str], key) -> tuple[str, bool] | None:
    """(id, exact) of the option `text` names: by name, by its key (company words dropped,
    Malay department words read in English), then fuzzy (difflib) or by a unique part."""
    if not text.strip() or not options:
        return None
    norm = _norm(text)
    for oid, name in options.items():
        if _norm(name) == norm:
            return oid, True
    want = key(text)
    keys = {oid: key(name) for oid, name in options.items()}
    for oid, k in keys.items():
        if k == want:
            return oid, True
    close = difflib.get_close_matches(want, list(keys.values()), n=1, cutoff=0.75)
    if close:
        return next(oid for oid, k in keys.items() if k == close[0]), False
    if len(want) >= 3:
        part = [oid for oid, k in keys.items() if want in k or (len(k) >= 3 and k in want)]
        if len(part) == 1:
            return part[0], False
    return None


def match_role(text: str) -> tuple[str, bool] | None:
    """(role key, exact) for what a sheet calls a role."""
    norm = _norm(text)
    if not norm:
        return None
    if norm in ROLE_LOOKUP:
        return ROLE_LOOKUP[norm], True
    padded = f" {norm} "
    for word in sorted(ROLE_LOOKUP, key=len, reverse=True):  # "pengurus cawangan" first
        if f" {word} " in padded:
            return ROLE_LOOKUP[word], False
    close = difflib.get_close_matches(norm, list(ROLE_LOOKUP), n=1, cutoff=0.8)
    return (ROLE_LOOKUP[close[0]], False) if close else None


def _status(notes: list[Note]) -> Literal["ok", "warn", "error"]:
    if any(n.level == "error" for n in notes):
        return "error"
    return "warn" if notes else "ok"


def _counts(rows: list[RowOut]) -> dict[str, int]:
    out = {"ok": 0, "warn": 0, "error": 0}
    for r in rows:
        out[r.status] += 1
    return out


def _error_text(e: HTTPException) -> str:
    detail: Any = e.detail  # ours carries {"code", "message"}
    msg = detail.get("message", "") if isinstance(detail, dict) else detail
    return render(msg, current_lang())


def _match_row(principal: Principal, org: Org, row: RowOut) -> list[Note]:
    """Fill the row's role, company and department from what the sheet said."""
    notes: list[Note] = []

    def note(field: str, level: Literal["warn", "error"], text: str) -> None:
        notes.append(Note(field=field, level=level, source="match", text=text))

    if row.role_text:
        found = match_role(row.role_text)
        if found is None:
            row.role = "staff"
            note(
                "role",
                "warn",
                tr(
                    'Role "{text}" is not one the app knows. Set to {role}.',
                    text=row.role_text,
                    role=i18n_labels.role_label("staff"),
                ),
            )
        else:
            row.role = found[0]
            if not found[1]:
                note(
                    "role",
                    "warn",
                    tr(
                        'Read "{text}" as {role}.',
                        text=row.role_text,
                        role=i18n_labels.role_label(found[0]),
                    ),
                )
    if row.role not in SCOPED_ROLES:  # workspace roles sit nowhere
        return notes

    sc = principal.scope
    branches = {b.id: b.name for b in org.branches.values()}
    if row.company_text:
        found = best_match(row.company_text, branches, _company_key)
        if found is None:
            note(
                "company",
                "error",
                tr('No company called "{text}". Pick one.', text=row.company_text),
            )
        else:
            row.branch_id = found[0]
            if not found[1]:
                note(
                    "company",
                    "warn",
                    tr(
                        'Matched "{text}" to {name}.',
                        text=row.company_text,
                        name=branches[found[0]],
                    ),
                )
    elif not sc.everything:
        row.branch_id = sc.branch_id
    elif len(branches) == 1:
        row.branch_id = next(iter(branches))

    if row.role == "branch_manager" or any(n.level == "error" for n in notes):
        return notes  # a branch manager runs the whole branch
    if row.department_text:
        pool = {
            d.id: d.name
            for d in org.departments.values()
            if row.branch_id is None or d.branch_id == row.branch_id
        }
        found = best_match(row.department_text, pool, _dept_key)
        if found is None:
            where = branches.get(row.branch_id or "")
            note(
                "department",
                "error",
                tr(
                    'No department called "{text}" in {company}.',
                    text=row.department_text,
                    company=where,
                )
                if where
                else tr('No department called "{text}". Pick one.', text=row.department_text),
            )
        elif (
            row.branch_id is None
            and sum(
                1
                for d in org.departments.values()
                if _dept_key(d.name) == _dept_key(pool[found[0]])
            )
            > 1
        ):
            note(
                "company",
                "error",
                tr(
                    'Several companies have a "{text}" department. Pick the company.',
                    text=pool[found[0]],
                ),
            )
        else:
            d = org.departments[found[0]]
            row.department_id, row.branch_id = d.id, d.branch_id
            if not found[1]:
                note(
                    "department",
                    "warn",
                    tr('Matched "{text}" to {name}.', text=row.department_text, name=d.name),
                )
    elif sc.kind == "department" and row.branch_id == sc.branch_id:
        row.department_id = sc.department_id
    return notes


async def _check_rows(
    db: AsyncSession, principal: Principal, org: Org, rows: list[RowOut], skip_place: set[int]
) -> None:
    """Every rule of adding one person, without adding anyone: a valid email and a name,
    once in the sheet, not already a member, a role this person may give, a place in their
    scope. `skip_place`: rows whose company or department could not be read (that error is
    already on the row). Adds "check" notes and sets each row's status."""
    emails = {r.email.strip().lower() for r in rows if r.email.strip()}
    users = {
        u.email: u for u in (await db.scalars(select(User).where(User.email.in_(emails)))).all()
    }
    members_here = (
        set(
            (
                await db.scalars(
                    select(Membership.user_id).where(
                        Membership.workspace_id == principal.workspace_id,
                        Membership.user_id.in_([u.id for u in users.values()]),
                    )
                )
            ).all()
        )
        if users
        else set()
    )
    seen: dict[str, int] = {}
    for r in rows:
        notes = [n for n in r.notes_found if n.source == "match"]

        def note(field: str, text: str, notes: list[Note] = notes) -> None:
            notes.append(Note(field=field, level="error", source="check", text=text))

        r.name = " ".join(r.name.split())
        r.email = r.email.strip().lower()
        r.existing_account = False
        if not r.name:
            note("name", tr("Name is missing."))
        if not r.email:
            note("email", tr("Email is missing."))
        else:
            try:
                _EMAIL.validate_python(r.email)
            except ValidationError:
                note("email", tr("{email} is not a valid email address.", email=r.email))
            else:
                if r.email in seen:
                    note(
                        "email",
                        tr(
                            "{email} is in the sheet twice (row {row}).",
                            email=r.email,
                            row=seen[r.email],
                        ),
                    )
                else:
                    seen[r.email] = r.row
                u = users.get(r.email)
                if u is not None and u.id in members_here:
                    note("member", tr("{email} is already a member.", email=r.email))
                elif u is not None:
                    r.existing_account = True
        if r.role not in ROLES:
            note("role", tr("Pick a role."))
        else:
            try:
                members._check_can_touch(principal, r.role, r.role)
            except HTTPException as e:
                note("role", _error_text(e))
            if r.row not in skip_place:
                try:
                    b, d = await members._placement(
                        db, principal, r.role, r.branch_id, r.department_id
                    )
                    r.branch_id, r.department_id = b, d
                except HTTPException as e:
                    note("place", _error_text(e))
        if r.role not in SCOPED_ROLES:
            r.branch_id = r.department_id = None
        r.branch_name = org.branches[r.branch_id].name if r.branch_id in org.branches else None
        r.department_name = (
            org.departments[r.department_id].name if r.department_id in org.departments else None
        )
        r.notes_found = notes
        r.status = _status(notes)


def _assignable(principal: Principal) -> list[str]:
    """The roles this person may give (as members._check_can_touch decides)."""
    if can(principal.role, "members.manage"):
        return [r for r in ROLES if r != "owner" or can(principal.role, "members.assign_owner")]
    allowed = members.TEAM_ROLES.get(principal.role, set())
    return [r for r in ROLES if r in allowed]


def _in_scope(principal: Principal, org: Org) -> Org:
    """Only the companies and departments this person may add people to."""
    sc = principal.scope
    if sc.everything:
        return org
    branches = [b for b in org.branches.values() if b.id == sc.branch_id]
    depts = [
        d
        for d in org.departments.values()
        if d.branch_id == sc.branch_id and (sc.kind != "department" or d.id == sc.department_id)
    ]
    return Org(branches, depts)


# What each scoped role does, for the template's list of roles.
ROLE_BLURBS = {
    "branch_manager": "Runs one company: its agents, work and people.",
    "hod": "Runs one department.",
    "supervisor": "Follows one department's work.",
    "staff": "Works with their own agents.",
}


# ---------------------------------------------------------------- endpoints


@router.get("/template")
async def template(
    principal: Principal = Depends(members.people_manager()),
    db: AsyncSession = Depends(get_db),
) -> Response:
    """An .xlsx to fill in: the columns with one made-up example row, and a second sheet of
    the companies, departments and roles this person may use (in their language)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    org = _in_scope(principal, await Org.load(db, principal.workspace_id))
    roles = _assignable(principal)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = tr("Members")[:31]
    heads = [tr("Name"), tr("Email"), tr("Role"), tr("Company"), tr("Department")]
    heads += [tr("Phone"), tr("Notes")]
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF9")
    for i, h in enumerate(heads, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = bold, fill
    branch = next(iter(org.branches.values()), None)
    dept = next((d for d in org.departments.values() if branch and d.branch_id == branch.id), None)
    role = "staff" if "staff" in roles else (roles[0] if roles else "staff")
    example = [
        "Siti Example",
        "siti@example.com",
        tr(i18n_labels.ROLES[role]),
        branch.name if branch else "",
        dept.name if dept else "",
        "012-345 6789",
        tr("Example row: replace it with your people."),
    ]
    for i, v in enumerate(example, 1):
        ws.cell(2, i, v).font = Font(italic=True, color="808080")
    for col, width in zip("ABCDEFG", (28, 32, 22, 30, 24, 18, 36), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    # The lists: company + department pairs, the roles, and each company once (dropdowns).
    lists = wb.create_sheet(tr("Lists")[:31])
    lists.append([tr("Company"), tr("Department"), "", tr("Role"), tr("What they can do"), "", ""])
    lists.cell(1, 7, tr("Companies"))
    for c in lists[1]:
        c.font = bold
    places: list[tuple[str, str]] = []
    for b in org.branches.values():
        ds = [d.name for d in org.departments.values() if d.branch_id == b.id]
        places += [(b.name, d) for d in ds] or [(b.name, "")]
    role_rows = [
        (
            tr(i18n_labels.ROLES[r]),
            tr(ROLE_BLURBS[r]) if r in ROLE_BLURBS else tr("Sees the whole workspace."),
        )
        for r in roles
    ]
    companies = [b.name for b in org.branches.values()]
    for i in range(max(len(places), len(role_rows), len(companies))):
        p = places[i] if i < len(places) else ("", "")
        rr = role_rows[i] if i < len(role_rows) else ("", "")
        co = companies[i] if i < len(companies) else ""
        lists.append([p[0], p[1], "", rr[0], rr[1], "", co])
    for col, width in zip("ABCDEFG", (30, 24, 4, 24, 44, 4, 30), strict=True):
        lists.column_dimensions[col].width = width
    sheet = lists.title.replace("'", "''")
    for column, letter, count in (("C", "D", len(role_rows)), ("D", "G", len(companies))):
        if count:
            dv = DataValidation(
                type="list", formula1=f"='{sheet}'!${letter}$2:${letter}${count + 1}"
            )
            dv.allow_blank = True
            ws.add_data_validation(dv)
            dv.add(f"{column}2:{column}{MAX_ROWS + 1}")
    out = io.BytesIO()
    wb.save(out)
    name = tr("members-template") + ".xlsx"
    return Response(
        out.getvalue(),
        media_type=XLSX,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )


@router.post("/preview")
async def preview(
    request: Request,
    name: str = Query(min_length=1, max_length=200),
    principal: Principal = Depends(members.people_manager()),
    db: AsyncSession = Depends(get_db),
) -> PreviewOut:
    data = bytearray()
    async for chunk in request.stream():
        data.extend(chunk)
        if len(data) > MAX_BYTES:
            raise api_error(
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                "file_too_large",
                "Files can be up to {mb} MB.",
                mb=MAX_BYTES // (1024 * 1024),
            )
    if not data:
        raise api_error(status.HTTP_400_BAD_REQUEST, "empty_file", "That file is empty.")
    sheet, rows = _read_rows(bytes(data), name)
    found = find_header(rows)
    mapped_by: Literal["headers", "ai"] = "headers"
    if found is None or "name" not in found[1] or "email" not in found[1]:
        ai = await _ai_columns(db, principal.workspace_id, rows) if rows else None
        if ai is None:
            raise api_error(
                status.HTTP_422_UNPROCESSABLE_CONTENT,
                "columns_not_found",
                "Could not find the Name and Email columns. Use the template, or name the "
                "columns Name and Email.",
            )
        found, mapped_by = ai, "ai"
    header, cols = found
    head_row = rows[header] if header >= 0 else []
    columns = [
        ColumnOut(
            field=f, header=head_row[i] if i < len(head_row) else "", column=i, letter=_letter(i)
        )
        for f, i in sorted(cols.items(), key=lambda kv: kv[1])
    ]

    org = await Org.load(db, principal.workspace_id)
    out: list[RowOut] = []
    capped = False
    for n, raw in enumerate(rows[header + 1 :], header + 2):
        if not any(raw):
            continue
        if len(out) >= MAX_ROWS:
            capped = True
            break

        def get(field: str, raw: list[str] = raw) -> str:
            i = cols.get(field)
            return raw[i] if i is not None and i < len(raw) else ""

        out.append(
            RowOut(
                row=n,
                name=get("name")[:200],
                email=get("email")[:320],
                phone=get("phone")[:60],
                notes=get("notes")[:500],
                role_text=get("role")[:80],
                company_text=get("company")[:200],
                department_text=get("department")[:200],
            )
        )
    if not out:
        raise api_error(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "no_rows",
            "The sheet has no people under its header row.",
        )
    skip: set[int] = set()
    for r in out:
        r.notes_found = _match_row(principal, org, r)
        if any(x.level == "error" and x.field in ("company", "department") for x in r.notes_found):
            skip.add(r.row)
    await _check_rows(db, principal, org, out, skip)
    return PreviewOut(
        file_name=name,
        sheet=sheet,
        header_row=header + 1,
        mapped_by=mapped_by,
        columns=columns,
        rows=out,
        counts=_counts(out),
        capped=capped,
    )


def _rows_out(rows: list[ImportRow]) -> list[RowOut]:
    return [RowOut(**r.model_dump()) for r in rows]


@router.post("/check")
async def check(
    body: RowsIn,
    principal: Principal = Depends(members.people_manager()),
    db: AsyncSession = Depends(get_db),
) -> CheckOut:
    """The rows again after the person fixed some by hand. Nothing is saved."""
    org = await Org.load(db, principal.workspace_id)
    rows = _rows_out(body.rows)
    await _check_rows(db, principal, org, rows, set())
    return CheckOut(rows=rows, counts=_counts(rows))


@router.post("")
async def import_members(
    body: RowsIn,
    principal: Principal = Depends(members.people_manager()),
    db: AsyncSession = Depends(get_db),
) -> ImportOut:
    """Add every row that passes the checks; skip (and say why) the rest. Each new account
    gets a one-time password, shown here only; existing accounts keep their own."""
    org = await Org.load(db, principal.workspace_id)
    rows = _rows_out(body.rows)
    await _check_rows(db, principal, org, rows, set())
    results: list[ResultOut] = []
    for r in rows:
        res = ResultOut(
            row=r.row,
            name=r.name,
            email=r.email,
            status="skipped",
            role=r.role,
            branch_name=r.branch_name,
            department_name=r.department_name,
        )
        results.append(res)
        if r.status == "error":
            res.message = " ".join(n.text for n in r.notes_found if n.level == "error")
            continue
        try:
            async with db.begin_nested():  # one bad row never undoes the others
                m, password = await members.add_one(
                    db,
                    principal,
                    email=r.email,
                    name=r.name,
                    role=r.role,
                    branch_id=r.branch_id,
                    department_id=r.department_id,
                    via="import",
                )
        except HTTPException as e:
            res.message = _error_text(e)
            continue
        res.status, res.user_id, res.temp_password = "added", m.user_id, password
        if password is None:
            res.message = tr("Already has an account. They sign in with their own password.")
    added = sum(1 for x in results if x.status == "added")
    await audit.record(
        db,
        principal.workspace_id,
        principal.actor,
        "members.imported",
        after={"added": added, "skipped": len(results) - added},
    )
    await db.commit()
    return ImportOut(results=results, added=added, skipped=len(results) - added)
