"""Adding members from an Excel or CSV sheet: the template, finding the header row and the
columns (English, Malay, the AI as a last resort), matching companies and departments,
the checks per row, a branch manager's limits, and the import with one-time passwords."""

import io

import httpx
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from agentic.api.routers import member_import
from agentic.core.db import SessionLocal
from agentic.engine import gateway
from agentic.models import AuditLog

from .conftest import csrf
from .test_agents import llm, office, temporal  # noqa: F401
from .test_office_roles import as_role, member, signed_in


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def preview(c: httpx.AsyncClient, data: bytes, name: str = "people.xlsx", **headers):
    return await c.post(
        "/api/members/import/preview",
        params={"name": name},
        content=data,
        headers={**csrf(c), "content-type": "application/octet-stream", **headers},
    )


def by_email(rows: list[dict]) -> dict[str, dict]:
    return {r["email"]: r for r in rows}


def texts(row: dict) -> str:
    return " | ".join(n["text"] for n in row["notes_found"])


# ---------------------------------------------------------------- reading the sheet


def test_header_words_in_english_and_malay():
    f = member_import.header_field
    assert f("Nama Penuh") == "name" and f("Full name") == "name"
    assert f("E-mel") == "email" and f("Emel") == "email" and f("Company email") == "email"
    assert f("Nama Syarikat") == "company" and f("Cawangan") == "company"
    assert f("Jabatan") == "department" and f("Dept.") == "department"
    assert f("Jawatan") == "role" and f("Peranan") == "role"
    assert f("No. Telefon") == "phone" and f("Bil") is None
    assert member_import.match_role("Kakitangan") == ("staff", True)
    assert member_import.match_role("Pengurus Cawangan") == ("branch_manager", True)
    assert member_import.match_role("ketua jabatan") == ("hod", True)
    assert member_import.match_role("Ketua Jabatan Kewangan") == ("hod", False)
    assert member_import.match_role("Supervisr") == ("supervisor", False)
    assert member_import.match_role("Kerani") is None


async def test_preview_finds_a_lower_header_row_and_matches_places(client):
    o = await office(client)
    data = xlsx(
        [
            ["Senarai Kakitangan 2026"],
            [],
            ["Bil", "Nama Penuh", "E-mel", "Jawatan", "Syarikat", "Jabatan", "No. Telefon"],
            [
                1,
                "Aminah Example",
                "Aminah@Example.com",
                "Kakitangan",
                "maju sdn bhd",
                "Kewangan",
                123456789,
            ],
            [2, "Badrul Example", "badrul@example.com", "Ketua Jabatan", "Maju", "Operasi", ""],
            [3, "Chong Example", "not-an-email", "Staf", "Maju", "", ""],
            [4, "Devi Example", "aminah@example.com", "", "Maju", "", ""],
            [5, "Eng Example", "eng@example.com", "Staff", "Unknown Holdings", "Finance", ""],
            [6, "Farid Example", "farid@example.com", "Kerani", "Maju Sdn. Bhd.", "Finanse", ""],
            [7, "Owner One", "owner@example.com", "Admin", "", "", ""],
            [8, "", "gina@example.com", "", "", "", ""],
            [],
        ]
    )
    r = await preview(client, data)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["header_row"] == 3 and p["mapped_by"] == "headers"
    assert {c["field"]: c["letter"] for c in p["columns"]} == {
        "name": "B",
        "email": "C",
        "role": "D",
        "company": "E",
        "department": "F",
        "phone": "G",
    }
    rows = p["rows"]
    assert [x["row"] for x in rows] == [4, 5, 6, 7, 8, 9, 10, 11]  # the blank row is skipped
    a = rows[0]
    assert a["email"] == "aminah@example.com" and a["status"] == "ok", texts(a)
    assert (a["role"], a["branch_name"], a["department_name"]) == (
        "staff",
        "Maju Sdn Bhd",
        "Finance",  # "Kewangan" read in English
    )
    assert a["phone"] == "123456789"
    b = rows[1]
    assert (b["role"], b["department_id"], b["status"]) == ("hod", o["depts"]["Operations"], "ok")
    assert rows[2]["status"] == "error" and "not a valid email" in texts(rows[2])
    assert rows[3]["status"] == "error" and "twice (row 4)" in texts(rows[3])
    assert rows[4]["status"] == "error" and 'No company called "Unknown Holdings"' in texts(rows[4])
    assert "needs a branch" not in texts(rows[4])  # one problem said once
    f = rows[5]
    assert f["status"] == "warn" and f["role"] == "staff"
    assert f["department_name"] == "Finance" and 'Matched "Finanse" to Finance' in texts(f)
    assert 'Role "Kerani"' in texts(f)
    assert rows[6]["status"] == "error" and "already a member" in texts(rows[6])
    assert rows[7]["status"] == "error" and "Name is missing" in texts(rows[7])
    assert p["counts"] == {"ok": 2, "warn": 1, "error": 5}

    # Nothing was saved.
    assert len((await client.get("/api/members")).json()) == 1

    # In Malay, the notes are in Malay.
    r = await preview(client, data, **{"x-lang": "ms"})
    assert "Tiada syarikat bernama" in texts(r.json()["rows"][4])


async def test_csv_with_semicolons_and_one_company_by_default(client):
    await office(client)
    data = "Name;Email;Role\nHana Example;hana@example.com;viewer\nIan Example;ian@example.com;\n"
    r = await preview(client, data.encode(), "people.csv")
    assert r.status_code == 200, r.text
    rows = by_email(r.json()["rows"])
    assert rows["hana@example.com"]["role"] == "viewer"
    assert rows["hana@example.com"]["branch_id"] is None  # workspace roles sit nowhere
    ian = rows["ian@example.com"]
    assert ian["role"] == "staff" and ian["branch_name"] == "Maju Sdn Bhd"  # the only company
    assert ian["status"] == "ok"


async def test_unclear_columns_ask_the_ai_and_say_so_when_it_cannot(client, monkeypatch):
    await office(client)
    # No header row at all: just people.
    data = xlsx(
        [["Joe Example", "joe@example.com", "Maju"], ["Kim Example", "kim@example.com", ""]]
    )

    async def down(*a, **kw):
        raise gateway.GatewayUnavailable("no model", [])

    monkeypatch.setattr(gateway, "chat", down)
    r = await preview(client, data)
    assert r.status_code == 422 and r.json()["code"] == "columns_not_found"

    class Reply:
        content = '{"header_row": -1, "columns": {"name": 0, "email": 1, "company": 2}}'

    async def answers(*a, **kw):
        assert kw["task"] == "members.import_columns"
        return Reply()

    monkeypatch.setattr(gateway, "chat", answers)
    r = await preview(client, data)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["mapped_by"] == "ai" and p["header_row"] == 0
    assert [x["email"] for x in p["rows"]] == ["joe@example.com", "kim@example.com"]
    assert p["rows"][0]["branch_name"] == "Maju Sdn Bhd"


async def test_bad_files(client):
    await office(client)
    r = await preview(client, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "old.xls")
    assert r.status_code == 400 and r.json()["code"] == "old_excel"
    r = await preview(client, b"PK\x03\x04broken", "broken.xlsx")
    assert r.status_code == 400 and r.json()["code"] == "bad_sheet"
    r = await preview(client, xlsx([["Name", "Email"]]))
    assert r.status_code == 422 and r.json()["code"] == "no_rows"
    # JSON only elsewhere: the raw upload is allowed on the preview path alone.
    r = await client.post(
        "/api/members/import",
        content=b"x",
        headers={**csrf(client), "content-type": "application/octet-stream"},
    )
    assert r.status_code == 415


# ---------------------------------------------------------------- who may add whom


async def test_a_branch_manager_adds_only_inside_their_branch(client):
    o = await office(client)
    b2 = (
        await client.post("/api/branches", json={"name": "Jaya Bhd"}, headers=csrf(client))
    ).json()
    a = o["branch"]["id"]
    bm = await as_role(client, "bm@example.com", "branch_manager", branch_id=a)
    try:
        data = xlsx(
            [
                ["Name", "Email", "Role", "Company", "Department"],
                ["Lee Example", "lee@example.com", "Staff", "", "Finance"],
                ["Mei Example", "mei@example.com", "Staff", "Jaya Bhd", ""],
                ["Nor Example", "nor@example.com", "Admin", "", ""],
                ["Oli Example", "oli@example.com", "Supervisor", "Maju", "Operations"],
            ]
        )
        r = await preview(bm, data)
        assert r.status_code == 200, r.text
        rows = by_email(r.json()["rows"])
        lee = rows["lee@example.com"]
        assert lee["status"] == "ok" and lee["branch_id"] == a  # their branch by default
        assert lee["department_id"] == o["depts"]["Finance"]
        assert rows["mei@example.com"]["status"] == "error"
        assert "only add people to" in texts(rows["mei@example.com"])
        assert rows["nor@example.com"]["status"] == "error"
        assert "only Head of department, Staff, Supervisor" in texts(rows["nor@example.com"])
        assert rows["oli@example.com"]["status"] == "ok"

        # The import checks again: a hand-edited row cannot reach another branch.
        sent = [
            {**lee, "branch_id": b2["id"], "department_id": None},
            rows["oli@example.com"],
            rows["nor@example.com"],
        ]
        r = await bm.post("/api/members/import", json={"rows": sent}, headers=csrf(bm))
        assert r.status_code == 200, r.text
        out = {x["email"]: x for x in r.json()["results"]}
        assert out["lee@example.com"]["status"] == "skipped"
        assert out["nor@example.com"]["status"] == "skipped"
        assert (
            out["oli@example.com"]["status"] == "added" and out["oli@example.com"]["temp_password"]
        )
        listed = {m["email"] for m in (await client.get("/api/members")).json()}
        assert "oli@example.com" in listed and "lee@example.com" not in listed

        # The template lists only their own branch and the roles they may give.
        r = await bm.get("/api/members/import/template")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        lists = [c for row in wb.worksheets[1].iter_rows(values_only=True) for c in row if c]
        assert "Maju Sdn Bhd" in lists and "Jaya Bhd" not in lists
        assert "Staff" in lists and "Admin" not in lists
    finally:
        await bm.aclose()


async def test_staff_cannot_import(client):
    o = await office(client)
    staff = await as_role(client, "s@example.com", "staff", branch_id=o["branch"]["id"])
    try:
        r = await staff.get("/api/members/import/template")
        assert r.status_code == 403
        r = await preview(staff, xlsx([["Name", "Email"], ["X", "x@example.com"]]))
        assert r.status_code == 403
    finally:
        await staff.aclose()


# ---------------------------------------------------------------- the import


async def test_fix_rows_then_import_with_one_time_passwords(client):
    o = await office(client)
    # An account that exists (it was a member before) keeps its own password.
    old = await member(client, "pat@example.com", "viewer")
    r = await client.delete(
        f"/api/members/{old['member']['user_id']}",
        headers={"content-type": "application/json", **csrf(client)},
    )
    assert r.status_code == 204

    data = xlsx(
        [
            ["Name", "Email", "Role", "Company", "Department"],
            ["Quinn Example", "quinn@example.com", "Staff", "Maju", "Finance"],
            ["Pat Example", "pat@example.com", "Viewer", "", ""],
            ["Rani Example", "rani@example.com", "Staff", "Nowhere Sdn Bhd", ""],
            ["Sam Example", "bad email", "Staff", "Maju", ""],
        ]
    )
    p = (await preview(client, data)).json()
    rows = by_email(p["rows"])
    assert rows["pat@example.com"]["existing_account"] is True
    rani = rows["rani@example.com"]
    assert rani["status"] == "error"

    # The person picks Rani's company by hand; the check says the row is fine now.
    rani = {**rani, "branch_id": o["branch"]["id"]}
    r = await client.post("/api/members/import/check", json={"rows": [rani]}, headers=csrf(client))
    assert r.status_code == 200, r.text
    assert r.json()["rows"][0]["status"] == "ok" and r.json()["counts"]["ok"] == 1

    sent = [rows["quinn@example.com"], rows["pat@example.com"], rani, rows["bad email"]]
    r = await client.post("/api/members/import", json={"rows": sent}, headers=csrf(client))
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["added"], body["skipped"]) == (3, 1)
    out = {x["email"]: x for x in body["results"]}
    quinn = out["quinn@example.com"]
    assert quinn["status"] == "added" and quinn["temp_password"]
    assert quinn["branch_name"] == "Maju Sdn Bhd" and quinn["department_name"] == "Finance"
    assert out["pat@example.com"]["status"] == "added"
    assert out["pat@example.com"]["temp_password"] is None  # an existing account
    assert out["rani@example.com"]["temp_password"]
    assert (
        out["bad email"]["status"] == "skipped"
        and "not a valid email" in out["bad email"]["message"]
    )

    # The one-time password signs them in (and they must pick their own).
    c = await signed_in("quinn@example.com", quinn["temp_password"])
    await c.aclose()
    members_now = {m["email"]: m for m in (await client.get("/api/members")).json()}
    assert members_now["quinn@example.com"]["role"] == "staff"
    assert members_now["quinn@example.com"]["department_id"] == o["depts"]["Finance"]

    # Importing the same sheet again adds nobody twice.
    r = await client.post("/api/members/import", json={"rows": sent[:1]}, headers=csrf(client))
    assert r.json()["results"][0]["status"] == "skipped"
    assert "already a member" in r.json()["results"][0]["message"]

    async with SessionLocal() as db:
        actions = (await db.scalars(select(AuditLog.action))).all()
        assert actions.count("members.imported") == 2
        added = (await db.scalars(select(AuditLog).where(AuditLog.action == "member.added"))).all()
        assert sum(1 for e in added if (e.after or {}).get("via") == "import") == 3


async def test_template_download(client):
    await office(client)
    r = await client.get("/api/members/import/template")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(member_import.XLSX)
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    first = list(wb.worksheets[0].iter_rows(values_only=True))
    assert first[0] == ("Name", "Email", "Role", "Company", "Department", "Phone", "Notes")
    assert first[1][1] == "siti@example.com" and first[1][3] == "Maju Sdn Bhd"
    lists = [c for row in wb.worksheets[1].iter_rows(values_only=True) for c in row if c]
    assert "Finance" in lists and "Branch manager" in lists and "Owner" in lists

    # Filled in and sent back as it is, the template reads cleanly.
    p = (await preview(client, r.content)).json()
    assert p["header_row"] == 1 and p["rows"][0]["status"] == "ok"

    r = await client.get("/api/members/import/template", headers={"x-lang": "ms"})
    wb = load_workbook(io.BytesIO(r.content))
    first = next(wb.worksheets[0].iter_rows(values_only=True))
    assert first[:5] == ("Nama", "E-mel", "Peranan", "Syarikat", "Jabatan")
    assert wb.worksheets[1].title == "Senarai"
    # The Malay sheet reads back too.
    p = (await preview(client, r.content)).json()
    assert p["rows"][0]["role"] == "staff" and p["rows"][0]["status"] == "ok"
