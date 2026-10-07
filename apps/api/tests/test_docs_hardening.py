"""P29: files, packs, forms, library and parsers hardening.

Packs only reference what the person may open; office roles keep files inside what they
manage; form blanks are copied instead of pulled away from their audience; agents keep out of
people's own files, other people's private assistants' work and other departments'
guidelines; private assistants' work is hidden even from admins; parsers refuse bombs;
formula-looking text is stored as text; the secrets scan sees more; the library index never
keeps stale passages; a wrong-size embedding model degrades to keywords.
"""

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, update

from agentic.core.db import SessionLocal
from agentic.documents import extract, render_xlsx, service
from agentic.models import Agent, DocFile, Document, KnowledgeChunk, Workspace

from .conftest import csrf
from .test_agents import llm, new_agent, office, temporal  # noqa: F401
from .test_documents import inline_reading, upload  # noqa: F401
from .test_library import REFUNDS, guideline, indexing, passages, wipe_chunks  # noqa: F401
from .test_office_roles import as_role


async def me(c) -> str:
    return (await c.get("/api/auth/me")).json()["user"]["id"]


async def private_assistant(
    client, o, owner_id: str, name: str = "Pa", dept: str = "Operations"
) -> Agent:
    """An agent turned into someone's private assistant."""
    a = await new_agent(client, o, name, dept)
    async with SessionLocal() as db:
        row = await db.get(Agent, a["id"])
        assert row is not None
        row.private, row.owner_user_id = True, owner_id
        await db.commit()
        await db.refresh(row)
        return row


async def made_by(agent: Agent, name: str, data: bytes, **extra) -> str:
    async with SessionLocal() as db:
        f = await service.create_file(
            db,
            workspace_id=agent.workspace_id,
            name=name,
            data=data,
            created_by=f"agent:{agent.id}",
            branch_id=agent.branch_id,
            agent_id=agent.id,
            source="generated",
            status="ready",
            **extra,
        )
        f.text = data.decode(errors="replace")
        await db.commit()
        return f.id


async def set_file(file_id: str, **values) -> None:
    async with SessionLocal() as db:
        await db.execute(update(DocFile).where(DocFile.id == file_id).values(**values))
        await db.commit()


# ================================================================ 1. packs


async def test_packs_reference_only_what_the_person_may_open(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
):
    o = await office(client)
    bid = o["branch"]["id"]
    bm = await as_role(client, "bm@example.com", "branch_manager", branch_id=bid)
    try:
        pa = await private_assistant(client, o, await me(bm))
        secret_out = await made_by(pa, "bank statement draft.txt", b"bank statement for bm")
        held = (await upload(client, "bank statement held.txt", b"x", branch_id=bid)).json()
        await set_file(held["id"], quarantined=True)
        ok = (await upload(client, "ssm certificate.txt", b"SSM cert", branch_id=bid)).json()

        async def new_pack(c, **item):
            return await c.post(
                "/api/packs",
                json={"title": "P", "branch_id": bid, "items": [{"label": "Item", **item}]},
                headers=csrf(c),
            )

        # The owner may not attach the manager's private assistant's work, nor a held file.
        assert (await new_pack(client, file_id=secret_out)).status_code == 400
        assert (await new_pack(client, file_id=held["id"])).status_code == 400
        # The manager may attach their own assistant's work; the owner then sees the pack
        # but not the file's name.
        r = await new_pack(bm, file_id=secret_out)
        assert r.status_code == 201, r.text
        theirs = r.json()
        assert theirs["items"][0]["file_name"] == "bank statement draft.txt"
        seen = (await client.get(f"/api/packs/{theirs['id']}")).json()
        assert seen["items"][0]["file_name"] is None

        # A pack for no company is a workspace role's: an office role gets their own company.
        r = await bm.post("/api/packs", json={"title": "Q", "items": []}, headers=csrf(bm))
        assert r.status_code == 201 and r.json()["branch_id"] == bid

        # Auto-match never picks the held-back file or someone's private assistant's work.
        r = await client.post(
            "/api/packs",
            json={"title": "R", "branch_id": bid, "items": [{"label": "Bank statement"}]},
            headers=csrf(client),
        )
        pack = r.json()
        m = (
            await client.post(f"/api/packs/{pack['id']}/auto-match", json={}, headers=csrf(client))
        ).json()
        assert m["pack"]["items"][0]["file_id"] not in (secret_out, held["id"])

        # Compiling checks again: a file held back after it was attached is left out.
        r = await new_pack(client, file_id=ok["id"])
        p2 = r.json()
        await set_file(ok["id"], quarantined=True)
        r = await client.post(f"/api/packs/{p2['id']}/compile", json={}, headers=csrf(client))
        assert r.status_code == 400  # nothing left to compile
    finally:
        await bm.aclose()


# ================================================================ 2. files PATCH


async def test_office_roles_keep_files_inside_what_they_manage(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
    indexing,  # noqa: F811
):
    o = await office(client)
    bid, fin, ops = o["branch"]["id"], o["depts"]["Finance"], o["depts"]["Operations"]
    b2 = (
        await client.post("/api/branches", json={"name": "Jaya Bhd"}, headers=csrf(client))
    ).json()
    hod = await as_role(client, "hod@example.com", "hod", branch_id=bid, department_id=fin)
    bm = await as_role(client, "bm@example.com", "branch_manager", branch_id=bid)
    try:
        f = (await upload(hod, "claims.txt", b"claims", branch_id=bid)).json()

        async def move(c, **body):
            return await c.patch(f"/api/files/{f['id']}", json=body, headers=csrf(c))

        assert (await move(hod, department_id=fin)).status_code == 200
        assert (await move(hod, department_id=None)).status_code == 403
        assert (await move(hod, department_id=ops)).status_code == 403
        assert (await move(hod, branch_id=None)).status_code == 403
        assert (await move(hod, name="claims 2.txt")).status_code == 200  # no move: fine
        assert (await move(bm, branch_id=None)).status_code == 403
        assert (await move(bm, branch_id=b2["id"])).status_code in (400, 403)
        assert (await move(bm, department_id=ops)).status_code == 200  # inside their company

        # Moving a guideline rebuilds its passages with the new audience.
        g = await guideline(client, "refunds.txt", REFUNDS, department_id=fin)
        r = await client.patch(
            f"/api/files/{g['id']}", json={"department_id": ops}, headers=csrf(client)
        )
        assert r.status_code == 200
        async with SessionLocal() as db:
            depts = set(
                (
                    await db.scalars(
                        select(KnowledgeChunk.department_id).where(
                            KnowledgeChunk.source_id == g["id"]
                        )
                    )
                ).all()
            )
        assert depts == {ops}
    finally:
        await hod.aclose()
        await bm.aclose()


# ================================================================ 3. forms


async def test_a_form_copies_a_guideline_it_does_not_own(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
    indexing,  # noqa: F811
):
    o = await office(client)
    bid = o["branch"]["id"]
    shared = await guideline(client, "claim form.txt", REFUNDS)  # workspace-wide
    bm = await as_role(client, "bm@example.com", "branch_manager", branch_id=bid)
    try:
        r = await bm.post(
            "/api/forms",
            json={"name": "Claims", "branch_id": bid, "file_id": shared["id"]},
            headers=csrf(bm),
        )
        assert r.status_code == 201, r.text
        blank = r.json()["template"]["id"]
        assert blank != shared["id"]
        async with SessionLocal() as db:
            orig = await db.get(DocFile, shared["id"])
            copy = await db.get(DocFile, blank)
            assert orig is not None and orig.branch_id is None and orig.library
            assert copy is not None and copy.branch_id == bid and copy.library
            chunk_branches = set(
                (
                    await db.scalars(
                        select(KnowledgeChunk.branch_id).where(KnowledgeChunk.source_id == blank)
                    )
                ).all()
            )
        assert chunk_branches == {bid}  # indexed at once, for its new audience
        assert await passages(shared["id"]) > 0  # the original keeps its passages
    finally:
        await bm.aclose()


# ================================================================ 4 + 5. agents, private work


async def test_agents_keep_out_of_personal_and_private_work(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
    indexing,  # noqa: F811
):
    from agentic.agents.doc_tools import _list_files, _read_file
    from agentic.agents.tools import ToolContext
    from agentic.search.viewer import for_agent

    o = await office(client)
    bid, fin = o["branch"]["id"], o["depts"]["Finance"]
    staff = await as_role(client, "siti@example.com", "staff", branch_id=bid)
    try:
        siti = await me(staff)
        desk = (
            await staff.post(
                "/api/desk/files",
                params={"name": "my payslip.txt", "branch_id": bid},
                content=b"payslip of siti",
                headers={**csrf(staff), "content-type": "application/octet-stream"},
            )
        ).json()
        company = (await upload(client, "price list.txt", b"price list", branch_id=bid)).json()
        fin_rule = await guideline(client, "finance rule.txt", "Finance rule", department_id=fin)
        ops = await new_agent(client, o, "Ops", "Operations")
        twin = await new_agent(client, o, "Siti twin", "Operations")
        pa = await private_assistant(client, o, siti, "Siti PA")
        pa_out = await made_by(pa, "pa notes.txt", b"private notes")
        doc = None
        async with SessionLocal() as db:
            t = await db.get(Agent, twin["id"])
            assert t is not None
            t.owner_user_id = siti
            d = Document(
                workspace_id=pa.workspace_id,
                branch_id=bid,
                agent_id=pa.id,
                title="PA letter",
                created_by=f"agent:{pa.id}",
            )
            db.add(d)
            await db.commit()
            doc = d.id

        async with SessionLocal() as db:
            ws = await db.get(Workspace, pa.workspace_id)
            for aid, sees in (
                (ops["id"], {company["id"]}),
                (twin["id"], {company["id"], desk["id"]}),
                (pa.id, {company["id"], desk["id"], pa_out}),
            ):
                a = await db.get(Agent, aid)
                assert a is not None
                ctx = ToolContext(db=db, agent=a, workspace=ws, task=None)
                for fid in (company["id"], desk["id"], pa_out, fin_rule["id"]):
                    out = await _read_file(ctx, {"file_id": fid})
                    assert ("no such file" not in out) == (fid in sees), (a.name, fid, out)
                listed = await _list_files(ctx, {})
                assert ("my payslip" in listed) == (desk["id"] in sees)
                # Document search follows the same rule.
                visible = set(
                    (await db.scalars(select(DocFile.id).where(for_agent(a).files()))).all()
                )
                assert (desk["id"] in visible) == (desk["id"] in sees)
                assert (pa_out in visible) == (pa_out in sees)
                assert fin_rule["id"] not in visible  # an Operations agent
                docs = set(
                    (await db.scalars(select(Document.id).where(for_agent(a).documents()))).all()
                )
                assert (doc in docs) == (a.id == pa.id)

        # P29 #5: the owner (a workspace role) does not see the assistant's work; Siti does.
        listed = {f["id"] for f in (await client.get("/api/files")).json()}
        assert pa_out not in listed and company["id"] in listed
        assert (await client.get(f"/api/files/{pa_out}")).status_code == 404
        assert (await client.get(f"/api/documents/{doc}")).status_code == 404
        assert (await staff.get(f"/api/files/{pa_out}")).status_code == 200
    finally:
        await staff.aclose()


# ================================================================ 6. parsers


def _bomb_docx(part_bytes: int = 4 * 1024 * 1024) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"\0" * part_bytes)
    return buf.getvalue()


def test_office_zip_bombs_are_refused_before_parsing():
    from agentic.forms import describe, fill

    out = extract.extract(_bomb_docx(), "x.docx")
    assert out.text == "" and "not read" in out.note
    with pytest.raises(extract.TooLarge):
        describe(_bomb_docx())
    with pytest.raises(extract.TooLarge):
        fill(_bomb_docx(), cells={"A1": 1})
    with pytest.raises(extract.TooLarge):
        extract.check_office_zip(_bomb_docx(1024 * 1024 * 2), max_unzipped=1024 * 1024)


def test_pictures_and_pages_are_capped():
    from PIL import Image

    strip = Image.new("L", (8000, 10))
    buf = io.BytesIO()
    strip.save(buf, format="PNG")
    out = extract._image(buf.getvalue())
    assert out.text == "" and "too large or too oddly shaped" in out.note
    assert extract.ocr_image(strip) == ""

    class Page:
        def __init__(self, w, h):
            self.size = (w, h)

        def get_size(self):
            return self.size

    a4 = extract._render_scale(Page(595, 842))
    assert a4 == extract.OCR_SCALE
    poster = extract._render_scale(Page(14400, 14400))
    assert (14400 * poster) ** 2 <= extract.OCR_MAX_PIXELS * 1.01


def test_unpack_report_keeps_at_most_200_skips():
    from agentic.intake import unpack

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(300):
            z.writestr(f"docs/empty{i}.txt", b"")
        z.writestr("docs/real.txt", b"hello")
    items = list(unpack.entries(buf.getvalue(), "u.zip"))
    report = items[-1]
    assert isinstance(report, unpack.Report)
    assert len(report.skipped) == unpack.MAX_SKIPPED and report.skipped_more == 100
    assert [e.name for e in items[:-1]] == ["real.txt"]
    # The upload's own buffer is read in place.
    again = list(unpack.entries(io.BytesIO(buf.getvalue()), "u.zip"))
    assert [e.name for e in again[:-1]] == ["real.txt"]


def test_form_sheets_are_scanned_within_a_cap():
    from agentic.forms import MAX_SCAN_ROWS, _scan

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(row=5000, column=1, value="far away")
    ws["A1"] = "Nama"
    rows = list(_scan(ws))
    assert len(rows) == MAX_SCAN_ROWS


# ================================================================ 7. formulas


def test_formula_looking_text_is_stored_as_text():
    from agentic.forms import fill

    data = render_xlsx.render(
        '| Item | Note |\n|---|---|\n| =HYPERLINK("http://x","y") | @SUM(A1) |\n| 5 | -3 |',
        [("Client", "=cmd|' /C calc'!A0")],
        "=1+1",
    )
    wb = load_workbook(io.BytesIO(data))
    det, table = wb["Details"], wb.worksheets[1]
    assert det["A1"].data_type == "s" and det["B3"].data_type == "s"
    assert table["A2"].data_type == "s" and table["B2"].data_type == "s"
    assert table["A3"].value == 5 and table["B3"].value == -3  # numbers stay numbers

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"], ws["C1"] = "Nama", "=SUM(1,2)"
    buf = io.BytesIO()
    wb.save(buf)
    out, _ = fill(buf.getvalue(), cells={"B5": "=1+1", "C1": "x"}, fields={"Nama": "+60123"})
    ws2 = load_workbook(io.BytesIO(out)).active
    assert ws2 is not None
    assert ws2["B5"].data_type == "s" and ws2["B5"].value == "=1+1"
    assert ws2["B1"].data_type == "s"
    assert ws2["C1"].data_type == "f"  # the form's own formula is kept


# ================================================================ 8. secrets scan


def _docx_with_header(header: str, body: str) -> bytes:
    from docx import Document as Docx

    d = Docx()
    d.add_paragraph(body)
    d.sections[0].header.paragraphs[0].text = header
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_scan_finds_new_key_kinds_and_office_parts():
    from agentic.intake import scan

    found = scan.scan("key: sk-ant-api03-" + "Ab1" * 10)
    assert "api_key" in found.record()["credentials"]
    url = "DATABASE_URL=postgres://app:S3cret_pw@db.internal:5432/app"
    found = scan.scan(url)
    assert "password" in found.record()["credentials"]
    assert "S3cret_pw" not in scan.mask(url, found)
    assert not scan.scan("postgres://user:password@localhost/db").any  # a placeholder
    data = _docx_with_header("Portal password: Xy7#kLm9q", "An ordinary letter.")
    assert "Xy7#kLm9q" in scan.office_text(data)
    assert "password" in scan.scan_bytes(data, "a.docx").record()["credentials"]


async def test_header_secrets_hold_uploads_and_generated_files_are_noted(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
):
    o = await office(client)
    bid = o["branch"]["id"]
    data = _docx_with_header("Portal password: Xy7#kLm9q", "An ordinary letter.")
    f = (await upload(client, "letter.docx", data, branch_id=bid)).json()
    got = (await client.get(f"/api/files/{f['id']}")).json()
    assert got["quarantined"] and "password" in got["sensitive"]["credentials"]

    agent = await new_agent(client, o, "Coder")
    async with SessionLocal() as db:
        a = await db.get(Agent, agent["id"])
        assert a is not None
    made = await made_by(
        a, "out.csv", b"user,password\nadmin,Pw: hunter2x9\nkey sk-ant-" + b"x1" * 15
    )
    async with SessionLocal() as db:
        g = await db.get(DocFile, made)
        assert g is not None and not g.quarantined
        assert "api_key" in g.sensitive["credentials"]


# ================================================================ 9. index race


async def test_index_file_drops_passages_when_the_file_leaves_the_library_meanwhile(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
    indexing,  # noqa: F811
    monkeypatch,
):
    from agentic.brain import embed
    from agentic.knowledge import indexer

    await office(client)
    g = await guideline(client, "refunds.txt", REFUNDS)
    assert await passages(g["id"]) > 0
    real = embed.embed

    async def slow_embed(texts):
        # Someone takes the file out of the library while its passages are being built.
        async with SessionLocal() as other:
            await other.execute(update(DocFile).where(DocFile.id == g["id"]).values(library=False))
            await other.commit()
        return await real(texts)

    monkeypatch.setattr(embed, "embed", slow_embed)
    async with SessionLocal() as db:
        assert await indexer.index_file(db, g["id"]) == 0
    assert await passages(g["id"]) == 0


# ================================================================ 10. embeddings


async def test_a_wrong_size_model_falls_back_to_keywords(
    client,
    llm,  # noqa: F811
    temporal,  # noqa: F811
    inline_reading,  # noqa: F811
    indexing,  # noqa: F811
    monkeypatch,
):
    from agentic.brain import embed
    from agentic.knowledge import indexer
    from agentic.knowledge import search as library

    await office(client)
    g = await guideline(client, "refunds.txt", REFUNDS)
    assert not indexer.dims_ok([0.0] * 10) and indexer.dims_ok(None)

    async def wrong(texts):
        return [[0.1] * 12 for _ in texts]

    monkeypatch.setattr(embed, "embed", wrong)
    async with SessionLocal() as db:
        assert await indexer.index_file(db, g["id"]) > 0  # stored without vectors
        n = await db.scalar(
            select(func.count())
            .select_from(KnowledgeChunk)
            .where(KnowledgeChunk.source_id == g["id"], KnowledgeChunk.embedding.is_not(None))
        )
        assert n == 0
        ws = (await db.get(DocFile, g["id"])).workspace_id  # type: ignore[union-attr]
        hits = await library.search(db, library.Reader(ws, everything=True), "damaged goods")
        assert hits and hits[0].source_id == g["id"]
