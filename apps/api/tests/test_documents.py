"""P10 Document Studio: files, company kits, templates, documents, exports and packs."""

import io
import json
from datetime import date

import httpx
import pytest
from sqlalchemy import select

from agentic.agents import dispatch, runtime
from agentic.core.db import SessionLocal
from agentic.documents import (
    blocks,
    checks,
    docx_template,
    extract,
    render_docx,
    render_pdf,
    render_xlsx,
)
from agentic.documents.fill import amount_in_words, fill
from agentic.models import DocFile, Document, Pack

from .conftest import csrf
from .test_agents import llm, messages, new_agent, new_task, office, temporal  # noqa: F401
from .test_office_roles import as_role

KIT = {
    "legal_name": "Maju Sdn Bhd",
    "reg_no": "201901012345 (1234567-X)",
    "address": "No 12, Jalan Teknologi 3\n47810 Petaling Jaya",
    "phone": "+60 3-1234 5678",
    "email": "hello@maju.my",
    "signatory_name": "Aina Rahman",
    "signatory_title": "Director",
    "currency": "RM",
    "tax_label": "SST",
    "tax_rate": 8,
    "payment_terms": "30 days",
}


@pytest.fixture
def inline_reading(monkeypatch):
    """Uploads are read right away instead of on the worker."""

    async def start_file_extract(file_id):
        async with SessionLocal() as db:
            from agentic.documents import service

            await service.process_file(db, file_id)

    monkeypatch.setattr(dispatch, "start_file_extract", start_file_extract)


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document as Docx

    d = Docx()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def upload(c: httpx.AsyncClient, name: str, data: bytes, **params) -> httpx.Response:
    return await c.post(
        "/api/files",
        params={"name": name, **params},
        content=data,
        headers={**csrf(c), "content-type": "application/octet-stream"},
    )


# ================================================================ pure functions


def test_fill_computes_items_words_and_drops_empty_optional_lines():
    body = (
        "**To:** {{client_name}}\nAttn: {{client_contact}}\n\n{{items}}\n\n"
        "In words: {{total_words}}\n\n- Notes: {{notes}}\n- Terms: {{company.payment_terms}}"
    )
    fields = [
        {"key": "client_name", "label": "Client name", "required": True},
        {"key": "client_contact", "label": "Attention to"},
        {"key": "notes", "label": "Notes"},
        {"key": "items", "label": "Line items", "type": "items", "required": True},
    ]
    values = {"items": [{"description": "Cleaning", "qty": 2, "unit_price": "1,250.50"}]}
    f = fill(body, values, fields, KIT, {"number": "QT-1"}, date(2026, 10, 2))
    assert f.totals == {"subtotal": 2501.0, "tax": 200.08, "total": 2701.08}
    assert "Attn:" not in f.markdown and "Notes" not in f.markdown  # optional and empty
    assert "[[Client name]]" in f.markdown and f.missing == ["Client name"]  # required
    assert "Terms: 30 days" in f.markdown
    assert "Ringgit Malaysia Two Thousand Seven Hundred One and Cents Eight Only" in f.markdown
    assert amount_in_words(1000) == "Ringgit Malaysia One Thousand Only"


def test_blocks_keep_line_breaks_and_tables():
    parsed = blocks.parse(
        "# Title\n\nLine one\nLine two\n\n| A | B |\n|---|---|\n| x | 1 |\n\n- a\n- b\n\n---"
    )
    assert [b.kind for b in parsed] == ["heading", "para", "table", "bullets", "rule"]
    assert parsed[1].lines == ["Line one", "Line two"]
    assert parsed[2].rows == [["x", "1"]]
    assert "Line one  \nLine two" in blocks.to_preview_markdown("Line one\nLine two")
    # A sign line stays a short line of underscores in the preview, not a divider.
    assert r"\_\_\_\_\_" in blocks.to_preview_markdown("Yours,\n_____\n**Aina**")


def test_checks_catch_bad_totals_marks_and_wrong_registration():
    md = (
        "Our reg 202001099999.\n\nTODO: confirm\n\n"
        "| No. | Item | Amount |\n|---|---|---|\n| 1 | A | 100.00 |\n| 2 | B | 50.00 |\n"
        "|  | **Subtotal** | 160.00 |\n|  | **Total** | 160.00 |"
    )
    f = fill(md, {}, [], KIT, {}, date(2026, 10, 2))
    found = checks.run(f, body=md, fields=[], values={}, kit=KIT, today=date(2026, 10, 2))
    texts = " | ".join(c["text"] for c in found)
    assert "Subtotal shows 160.00 but the amounts above add up to 150.00" in texts
    assert "TODO" in texts
    assert "202001099999 does not match the company kit" in texts
    good = fill("Fine.", {}, [], KIT, {}, date(2026, 10, 2))
    assert checks.run(good, body="Fine.", fields=[], values={}, kit=KIT, today=date.today()) == []


def test_pdf_word_and_excel_exports_are_real_files():
    from docx import Document as Docx
    from openpyxl import load_workbook
    from pypdf import PdfReader

    md = "# Quote\n\nHello **there**.\n\n| Item | Amount |\n|---|---|\n| A | 1,250.00 |"
    lh = render_pdf.Letterhead.from_kit(KIT)
    pdf = render_pdf.render(md, "Quote", lh)
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "Maju Sdn Bhd" in text and "Quote" in text and "Page 1 of 1" in text
    word = Docx(io.BytesIO(render_docx.render(md, "Quote", lh)))
    assert any("Hello" in p.text for p in word.paragraphs) and len(word.tables) == 1
    wb = load_workbook(io.BytesIO(render_xlsx.render(md, [("Number", "QT-1")], "Quote")))
    assert wb.sheetnames == ["Details", "Quote"] and wb["Quote"]["B2"].value == 1250


def test_word_templates_are_filled_in_place_with_item_rows():
    data = _docx(
        ["Dear {{client_name}},", "Total: {{total}}"],
        [["Item", "Qty", "Amount"], ["{{item.description}}", "{{item.qty}}", "{{item.amount}}"]],
    )
    assert docx_template.scan(data) == ["client_name", "total", "items"]
    from agentic.documents.fill import context

    ctx, info = context(
        {"client_name": "Bina", "items": [{"description": "A", "qty": 2, "unit_price": 10}]},
        [{"key": "items", "type": "items"}],
        KIT,
        {},
        date(2026, 10, 2),
    )
    out, missing = docx_template.fill(data, ctx, info.items, [])
    text = docx_template.text_of(out)
    assert "Dear Bina," in text and "| A | 2 | 20.00 |" in text and missing == []


def test_extract_reads_word_excel_csv_and_pdf():
    from openpyxl import Workbook

    assert "Hello Word" in extract.extract(_docx(["Hello Word"]), "a.docx").text
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Name", "Amount"])
    ws.append(["Bina", 12])
    buf = io.BytesIO()
    wb.save(buf)
    assert "| Bina | 12 |" in extract.extract(buf.getvalue(), "b.xlsx").text
    assert "| x | 1 |" in extract.extract(b"a,b\nx,1\n", "c.csv").text
    pdf = render_pdf.render("Certificate of incorporation\n\nMaju Sdn Bhd", "Cert")
    out = extract.extract(pdf, "cert.pdf")
    assert out.pages == 1 and "[page 1]" in out.text and "Maju Sdn Bhd" in out.text


# ================================================================ files API


async def test_upload_read_list_download_and_guard(client, llm, temporal, inline_reading):
    o = await office(client)
    await client.put(
        "/api/ai/groups/fast",
        json={"members": [{"provider_id": o["provider"]["id"], "model_id": "m1"}]},
        headers=csrf(client),
    )
    llm.say(
        json.dumps(
            {
                "kind": "SSM certificate",
                "title": "Certificate of incorporation",
                "summary": "Maju Sdn Bhd was incorporated in 2019.",
                "fields": {"Registration no.": "201901012345"},
                "expires_on": "2020-01-01",
            }
        )
    )
    pdf = render_pdf.render(
        "Certificate of incorporation\n\nMaju Sdn Bhd 201901012345\n\nValid until 1 January 2020",
        "Cert",
    )
    r = await upload(client, "ssm.pdf", pdf, branch_id=o["branch"]["id"])
    assert r.status_code == 201, r.text
    f = r.json()
    assert f["status"] == "ready" and f["kind"] == "SSM certificate" and f["expired"] is True
    assert f["fields"]["Registration no."] == "201901012345"

    detail = (await client.get(f"/api/files/{f['id']}")).json()
    assert "Maju Sdn Bhd" in detail["text"]
    listed = (await client.get("/api/files", params={"q": "ssm"})).json()
    assert [x["id"] for x in listed] == [f["id"]]
    dl = await client.get(f"/api/files/{f['id']}/download")
    assert dl.content == pdf and "attachment" in dl.headers["content-disposition"]
    assert "sandbox" in dl.headers["content-security-policy"]

    # Raw bodies only on the upload path, and the CSRF token is still required.
    bad = await client.post(
        "/api/documents",
        content=b"x",
        headers={**csrf(client), "content-type": "application/octet-stream"},
    )
    assert bad.status_code == 415
    no_token = await client.post(
        "/api/files",
        params={"name": "a.txt"},
        content=b"hi",
        headers={"content-type": "application/octet-stream"},
    )
    assert no_token.status_code == 403
    empty = await upload(client, "a.txt", b"")
    assert empty.status_code == 400

    # A report's last date is not an expiry: kept only when the text talks about validity.
    llm.say(
        json.dumps(
            {"kind": "CSV file", "title": "Ad spend", "summary": "x", "expires_on": "2020-09-30"}
        )
    )
    rep = (await upload(client, "ads.csv", b"date,spend\n2020-09-30,10\n")).json()
    assert rep["expires_on"] is None and rep["expired"] is False


async def test_office_roles_only_see_their_files(client, llm, temporal, inline_reading):
    o = await office(client)
    mine = (await upload(client, "owner.txt", b"owner notes")).json()
    staff = await as_role(client, "staff@example.com", "staff", branch_id=o["branch"]["id"])
    theirs = (await upload(staff, "staff.txt", b"staff notes")).json()
    assert [f["id"] for f in (await staff.get("/api/files")).json()] == [theirs["id"]]
    assert (await staff.get(f"/api/files/{mine['id']}")).status_code == 404
    seen = {f["id"] for f in (await client.get("/api/files")).json()}
    assert seen == {mine["id"], theirs["id"]}


# ================================================================ kits, templates, documents


async def test_kit_templates_document_lifecycle_and_exports(client, llm, temporal):
    o = await office(client)
    bid = o["branch"]["id"]
    r = await client.put(
        f"/api/company-kits/{bid}", json={"data": {**KIT, "bogus": "x"}}, headers=csrf(client)
    )
    assert r.status_code == 200 and "bogus" not in r.json()["data"]
    assert (r.json()["filled"], r.json()["total"]) == (11, 18)  # no bank, tax no., website...

    tpls = (await client.get("/api/doc-templates")).json()
    assert {"Quotation", "Invoice", "Official letter", "Meeting minutes"} <= {
        t["name"] for t in tpls
    }
    quote = next(t for t in tpls if t["name"] == "Quotation")

    r = await client.post(
        "/api/documents",
        json={"template_id": quote["id"], "branch_id": bid, "title": "Quote for Bina"},
        headers=csrf(client),
    )
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["number"] == f"QT-{date.today().year}-0001" and doc["status"] == "draft"
    assert doc["errors"] >= 2  # client name, subject and items are required

    # Approval refuses while the checks find errors.
    r = await client.post(
        f"/api/documents/{doc['id']}/status", json={"status": "approved"}, headers=csrf(client)
    )
    assert r.status_code == 409 and r.json()["code"] == "checks_failed"

    values = {
        "client_name": "Syarikat Bina Sdn Bhd",
        "subject": "Office cleaning",
        "items": [{"description": "Monthly cleaning", "qty": 12, "unit_price": 1850}],
    }
    r = await client.patch(
        f"/api/documents/{doc['id']}", json={"values": values}, headers=csrf(client)
    )
    doc = r.json()
    assert doc["version"] == 2 and doc["errors"] == 0, doc["checks"]
    assert doc["totals"]["total"] == 23976.0 and "Syarikat Bina" in doc["preview"]

    # Unsaved preview changes nothing.
    p = await client.post(
        f"/api/documents/{doc['id']}/preview",
        json={"values": {**values, "client_name": ""}},
        headers=csrf(client),
    )
    assert "Client name" in p.json()["missing"]
    assert (await client.get(f"/api/documents/{doc['id']}")).json()["values"]["client_name"]

    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK"), ("xlsx", b"PK")):
        e = await client.get(f"/api/documents/{doc['id']}/export", params={"format": fmt})
        assert e.status_code == 200 and e.content.startswith(magic), fmt

    r = await client.post(
        f"/api/documents/{doc['id']}/status", json={"status": "approved"}, headers=csrf(client)
    )
    assert r.json()["status"] == "approved" and r.json()["approved_by"]
    locked = await client.patch(
        f"/api/documents/{doc['id']}", json={"title": "x"}, headers=csrf(client)
    )
    assert locked.status_code == 409
    await client.post(
        f"/api/documents/{doc['id']}/status", json={"status": "draft"}, headers=csrf(client)
    )
    versions = (await client.get(f"/api/documents/{doc['id']}/versions")).json()
    assert [v["version"] for v in versions] == [1]
    r = await client.post(
        f"/api/documents/{doc['id']}/restore", json={"version": 1}, headers=csrf(client)
    )
    assert r.json()["values"] == {} and r.json()["version"] == 3

    second = (
        await client.post(
            "/api/documents",
            json={"template_id": quote["id"], "branch_id": bid},
            headers=csrf(client),
        )
    ).json()
    assert second["number"].endswith("-0002")


async def test_word_template_upload_becomes_a_template(client, llm, temporal, inline_reading):
    await office(client)
    data = _docx(["Dear {{client_name}},", "Ref {{doc.number}}"])
    f = (await upload(client, "letter.docx", data)).json()
    r = await client.post(
        "/api/doc-templates/from-docx",
        json={"file_id": f["id"], "name": "Our letter"},
        headers=csrf(client),
    )
    assert r.status_code == 201, r.text
    assert [x["key"] for x in r.json()["fields"]] == ["client_name"]
    d = (
        await client.post(
            "/api/documents",
            json={"template_id": r.json()["id"], "values": {"client_name": "Bina"}},
            headers=csrf(client),
        )
    ).json()
    assert d["word_template"] is True and "Dear Bina," in d["preview"]
    e = await client.get(f"/api/documents/{d['id']}/export", params={"format": "docx"})
    assert "Dear Bina," in docx_template.text_of(e.content)
    plain = (await upload(client, "plain.docx", _docx(["No fields here"]))).json()
    r = await client.post(
        "/api/doc-templates/from-docx",
        json={"file_id": plain["id"], "name": "Nope"},
        headers=csrf(client),
    )
    assert r.json()["code"] == "no_placeholders"


async def test_ai_helpers_rewrite_fill_and_review(client, llm, temporal):
    o = await office(client)
    for g in ("fast",):
        await client.put(
            f"/api/ai/groups/{g}",
            json={"members": [{"provider_id": o["provider"]["id"], "model_id": "m1"}]},
            headers=csrf(client),
        )
    tpls = (await client.get("/api/doc-templates")).json()
    inv = next(t for t in tpls if t["name"] == "Invoice")
    doc = (
        await client.post(
            "/api/documents",
            json={"template_id": inv["id"], "branch_id": o["branch"]["id"]},
            headers=csrf(client),
        )
    ).json()
    llm.say("We kindly request payment.")
    r = await client.post(
        f"/api/documents/{doc['id']}/rewrite",
        json={"text": "Pay now.", "instruction": "more polite"},
        headers=csrf(client),
    )
    assert r.json()["text"] == "We kindly request payment."
    llm.say(json.dumps({"client_name": "Bina", "due_date": "2026-11-01", "made_up": "x"}))
    r = await client.post(
        f"/api/documents/{doc['id']}/ai-fill",
        json={"request": "Invoice Bina, due 1 Nov"},
        headers=csrf(client),
    )
    assert r.json()["values"] == {"client_name": "Bina", "due_date": "2026-11-01"}
    llm.say(json.dumps({"issues": [{"level": "warn", "text": "No reference."}]}))
    r = await client.post(f"/api/documents/{doc['id']}/review", json={}, headers=csrf(client))
    assert r.status_code == 200, r.text
    assert r.json()["issues"] == [{"level": "warn", "text": "No reference."}]


# ================================================================ packs


async def test_pack_matches_files_and_compiles_one_pdf(client, llm, temporal, inline_reading):
    from pypdf import PdfReader

    o = await office(client)
    bid = o["branch"]["id"]
    await client.put(f"/api/company-kits/{bid}", json={"data": KIT}, headers=csrf(client))
    cert = render_pdf.render("Certificate of incorporation", "SSM")
    f1 = (await upload(client, "company registration certificate.pdf", cert, branch_id=bid)).json()
    f2 = (
        await upload(
            client, "bank statement september.csv", b"date,amount\n1/9,100\n", branch_id=bid
        )
    ).json()
    tpls = (await client.get("/api/doc-templates")).json()
    cover = next(t for t in tpls if t["name"].startswith("Cover letter"))
    letter = (
        await client.post(
            "/api/documents",
            json={
                "template_id": cover["id"],
                "branch_id": bid,
                "values": {"client_name": "Agency", "submission": "Cleaning", "enclosures": "SSM"},
            },
            headers=csrf(client),
        )
    ).json()
    r = await client.post(
        "/api/packs",
        json={
            "title": "Cleaning tender pack",
            "branch_id": bid,
            "items": [
                {"label": "Company registration certificate"},
                {
                    "label": "Latest 3 months bank statements",
                    # A long AI-written hint must not drown the label.
                    "hint": "Official bank-stamped or certified true copies showing company "
                    "account name, account number, and recent transactions.",
                },
                {"label": "Cover letter", "document_id": letter["id"]},
                {"label": "Audited accounts", "required": False},
            ],
        },
        headers=csrf(client),
    )
    assert r.status_code == 201, r.text
    pack = r.json()
    assert pack["progress"] == {"total": 4, "required": 3, "ready": 1, "missing": 2}

    r = await client.post(f"/api/packs/{pack['id']}/auto-match", json={}, headers=csrf(client))
    assert r.status_code == 200, r.text
    m = r.json()
    assert m["matched"] == 2
    by = {i["label"]: i for i in m["pack"]["items"]}
    assert by["Company registration certificate"]["file_id"] == f1["id"]
    bank = by["Latest 3 months bank statements"]
    assert bank["file_id"] == f2["id"] and bank["auto"]

    r = await client.post(f"/api/packs/{pack['id']}/compile", json={}, headers=csrf(client))
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["skipped"] == [] and out["pack"]["status"] == "compiled"
    data = (await client.get(f"/api/files/{out['file_id']}/download")).content
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) == out["pages"] == 4  # cover + three items
    first = reader.pages[0].extract_text()
    assert "Cleaning tender pack" in first and "Contents" in first and "Cover letter" in first
    assert "Page 4 of 4" in reader.pages[3].extract_text()

    # Changing the checklist marks the compiled PDF out of date.
    items = [*m["pack"]["items"], {"label": "Extra"}]
    r = await client.patch(f"/api/packs/{pack['id']}", json={"items": items}, headers=csrf(client))
    assert r.json()["status"] == "collecting" and r.json()["progress"]["missing"] == 1


# ================================================================ agents


async def test_agent_reads_files_and_drafts_a_checked_document(
    client, llm, temporal, inline_reading
):
    o = await office(client)
    bid = o["branch"]["id"]
    await client.put(f"/api/company-kits/{bid}", json={"data": KIT}, headers=csrf(client))
    agent = await new_agent(client, o, "Aina")
    f = (
        await upload(
            client, "enquiry.txt", b"Bina wants 12 months cleaning at RM1850", branch_id=bid
        )
    ).json()
    r = await client.post(
        "/api/tasks",
        json={"title": "Quote Bina", "assignee_agent_id": agent["id"], "file_ids": [f["id"]]},
        headers=csrf(client),
    )
    task = r.json()
    llm.call("read_file", file_id=f["id"]).call(
        "draft_document",
        template="quotation",
        title="Quotation for Bina",
        values={
            "client_name": "Syarikat Bina",
            "subject": "Cleaning",
            "items": [{"description": "Monthly cleaning", "qty": 12, "unit_price": 1850}],
        },
    ).say("Drafted the quotation for review.")
    assert (await runtime.run_task_step(task["id"])).state == "done"
    msgs = await messages(task["id"])
    assert "Files attached to this task" in msgs[0].content and f["id"] in msgs[0].content
    tool_out = [m.content for m in msgs if m.role == "tool"]
    assert "12 months cleaning" in tool_out[0]
    assert "Document drafted" in tool_out[1] and "Checks: all clear." in tool_out[1]
    async with SessionLocal() as db:
        d = await db.scalar(select(Document).where(Document.task_id == task["id"]))
        assert d is not None and d.status == "review" and d.agent_id == agent["id"]
        assert d.number.startswith("QT-")


async def test_agent_tools_stay_inside_their_company_and_fill_packs(
    client, llm, temporal, inline_reading
):
    from agentic.agents.doc_tools import _pack_attach, _pack_status, _read_file
    from agentic.agents.tools import ToolContext
    from agentic.models import Agent, Workspace

    o = await office(client)
    b2 = (
        await client.post("/api/branches", json={"name": "Jaya Bhd"}, headers=csrf(client))
    ).json()
    agent = await new_agent(client, o, "Aina")
    other = (await upload(client, "jaya secret.txt", b"secret", branch_id=b2["id"])).json()
    ours = (
        await upload(
            client, "ssm.txt", b"[page 1]\nfirst\n\n[page 2]\nsecond", branch_id=o["branch"]["id"]
        )
    ).json()
    pack = (
        await client.post(
            "/api/packs",
            json={"title": "P", "branch_id": o["branch"]["id"], "items": [{"label": "SSM"}]},
            headers=csrf(client),
        )
    ).json()
    async with SessionLocal() as db:
        a = await db.get(Agent, agent["id"])
        ws = await db.get(Workspace, a.workspace_id)
        ctx = ToolContext(db=db, agent=a, workspace=ws, task=None)
        assert "no such file" in await _read_file(ctx, {"file_id": other["id"]})
        page2 = await _read_file(ctx, {"file_id": ours["id"], "pages": "2"})
        assert "second" in page2 and "first" not in page2
        item = pack["items"][0]["id"]
        out = await _pack_attach(
            ctx, {"pack_id": pack["id"], "item_id": item, "file_id": ours["id"]}
        )
        assert "1 ready, 0 missing" in out
        assert "✓" in await _pack_status(ctx, {"pack_id": pack["id"]})
        p = await db.get(Pack, pack["id"])
        assert p is not None and p.items[0]["file_id"] == ours["id"]
        f = await db.get(DocFile, other["id"])
        assert f is not None
