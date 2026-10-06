"""P27: company forms people fill in and hand back by a deadline.

When rounds open and close (a window that wraps the month end, a missed round still shown as
late), reading and filling an Excel form (labels, table rows, cells, merged cells, formulas
kept), ready-made forms, who sees and manages them, handing in with attachments, the
manager's round view (handed in, missing), returning and accepting, and the AI filling a
form into the person's draft.
"""

import io
from datetime import date

from openpyxl import load_workbook

from agentic.agents import runtime
from agentic.core.db import SessionLocal
from agentic.forms import describe, fill, missed, state, window
from agentic.forms.starters import STARTERS, build
from agentic.models import FormSubmission

from .conftest import csrf
from .test_agents import llm, office, temporal  # noqa: F401
from .test_office_roles import as_role
from .test_twins import ANSWERS

MONTHLY = {"every": "month", "from_day": 1, "to_day": 15}
WRAPS = {"every": "month", "from_day": 30, "to_day": 3}


def test_rounds_open_close_wrap_and_show_late():
    w = window(MONTHLY, date(2026, 10, 6))
    assert (w.period, w.opens, w.due) == ("2026-10", date(2026, 10, 1), date(2026, 10, 15))
    assert window(MONTHLY, date(2026, 10, 20)).period == "2026-11"  # type: ignore[union-attr]
    w = window(WRAPS, date(2026, 10, 2))  # inside September's round (30 Sep - 3 Oct)
    assert (w.period, w.due) == ("2026-09", date(2026, 10, 3))  # type: ignore[union-attr]
    assert missed(WRAPS, date(2026, 10, 6)).period == "2026-09"  # type: ignore[union-attr]
    assert missed(MONTHLY, date(2026, 10, 6)) is None
    year = window({"every": "year", "month": 12, "from_day": 1, "to_day": 31}, date(2026, 6, 1))
    assert (year.period, year.opens) == ("2026", date(2026, 12, 1))  # type: ignore[union-attr]
    assert window({"every": "none"}, date(2026, 6, 1)) is None
    assert state(window(MONTHLY, date(2026, 10, 14)), date(2026, 10, 14), None) == "due_soon"
    assert state(window(MONTHLY, date(2026, 10, 5)), date(2026, 10, 5), None) == "open"
    assert state(window(MONTHLY, date(2026, 10, 5)), date(2026, 10, 5), "submitted") == "submitted"
    assert state(missed(WRAPS, date(2026, 10, 6)), date(2026, 10, 6), None) == "late"


def test_ready_made_forms_read_and_fill_in_both_languages():
    for s in STARTERS:
        for lang in ("en", "ms"):
            _, data = build(s.key, lang, "Contoh Sdn Bhd")
            assert load_workbook(io.BytesIO(data)).worksheets[0]["A1"].value == "CONTOH SDN BHD"
    _, data = build("expense_claim", "ms", "Contoh Sdn Bhd")
    layout = describe(data)
    assert "A4 'Nama'" in layout and "C7 'Perkara'" in layout and "=SUM(E8:E22)" in layout
    out, report = fill(
        data,
        fields={"Nama": "Aina binti Ali", "Bulan": "Oktober 2026", "Tiada": "x"},
        rows=[
            {"Tarikh": "2026-10-03", "Perkara": "Grab ke PPD", "Jumlah (RM)": "RM23.50"},
            {"Perkara": "Kertas A4", "Jumlah (RM)": 45},
        ],
        cells={"F8": "resit hilang", "E23": 999},
    )
    ws = load_workbook(io.BytesIO(out)).worksheets[0]
    assert ws["C4"].value == "Aina binti Ali"  # the merged blank right of the label
    assert ws["C8"].value == "Grab ke PPD" and ws["E8"].value == 23.5 and ws["E9"].value == 45
    assert ws["F8"].value == "resit hilang"
    assert ws["E23"].value == "=SUM(E8:E22)"  # the total stays a formula
    assert "no label 'Tiada' in the form" in report and "kept the formula in E23" in report


async def test_forms_hand_in_review_and_who_sees_what(client, llm, temporal):
    o = await office(client)
    b = o["branch"]["id"]
    staff = await as_role(
        client, "siti@example.com", "staff", branch_id=b, department_id=o["depts"]["Finance"]
    )
    other = (
        await client.post("/api/branches", json={"name": "Jaya Bhd"}, headers=csrf(client))
    ).json()
    far = await as_role(client, "lim@example.com", "staff", branch_id=other["id"])
    try:
        starters = (await client.get("/api/forms/starters")).json()
        assert {"expense_claim", "travel_claim", "leave"} <= {s["key"] for s in starters}
        r = await client.post(
            "/api/forms/starters",
            json={"key": "expense_claim", "branch_id": b},
            headers=csrf(client),
        )
        assert r.status_code == 201, r.text
        form = r.json()
        assert form["template"]["name"].endswith(".xlsx") and form["can_manage"]
        assert form["schedule"]["every"] == "month" and form["progress"]["expected"] >= 1

        mine = (await staff.get("/api/forms")).json()
        assert [f["id"] for f in mine] == [form["id"]] and mine[0]["can_manage"] is False
        assert mine[0]["state"] in ("open", "due_soon", "upcoming", "late")
        assert (await far.get("/api/forms")).json() == []
        assert (await far.get(f"/api/forms/{form['id']}/download")).status_code == 404
        dl = await staff.get(f"/api/forms/{form['id']}/download")
        assert dl.status_code == 200 and dl.content[:2] == b"PK"
        # Staff do not add forms.
        r = await staff.post(
            "/api/forms/starters", json={"key": "leave", "branch_id": b}, headers=csrf(staff)
        )
        assert r.status_code == 403

        # Hand in: the filled form and a receipt, uploaded to the person's own workspace.
        up = []
        for name, body in (("tuntutan.xlsx", dl.content), ("resit.txt", b"Resit Grab RM23.50")):
            r = await staff.post(
                "/api/desk/files",
                params={"name": name},
                content=body,
                headers={**csrf(staff), "content-type": "application/octet-stream"},
            )
            assert r.status_code == 201, r.text
            up.append(r.json()["id"])
        r = await staff.post(
            f"/api/forms/{form['id']}/submit",
            json={"file_ids": up, "note": "Bulan ini"},
            headers=csrf(staff),
        )
        assert r.status_code == 200, r.text
        assert r.json()["state"] == "submitted" and len(r.json()["mine"]["files"]) == 2

        # The manager sees who handed in, returns it with a note; the person sees why.
        rounds = (await client.get(f"/api/forms/{form['id']}/submissions")).json()
        row = next(p for p in rounds["people"] if p["name"] == "Siti")
        assert row["state"] == "submitted" and len(row["submission"]["files"]) == 2
        sid = row["submission"]["id"]
        r = await client.post(
            f"/api/form-submissions/{sid}/review", json={"decision": "return"}, headers=csrf(client)
        )
        assert r.status_code == 400  # say what to fix
        r = await client.post(
            f"/api/form-submissions/{sid}/review",
            json={"decision": "return", "note": "Resit kedua tiada."},
            headers=csrf(client),
        )
        assert r.json()["status"] == "returned"
        back = (await staff.get("/api/forms")).json()[0]
        assert back["state"] == "returned" and back["mine"]["review_note"] == "Resit kedua tiada."
        r = await staff.post(
            f"/api/forms/{form['id']}/submit", json={"file_ids": up}, headers=csrf(staff)
        )
        assert r.json()["state"] == "submitted"
        r = await staff.post(
            f"/api/form-submissions/{sid}/review", json={"decision": "accept"}, headers=csrf(staff)
        )
        assert r.status_code == 403  # only managers review
        r = await client.post(
            f"/api/form-submissions/{sid}/review", json={"decision": "accept"}, headers=csrf(client)
        )
        assert r.json()["status"] == "accepted"
        r = await staff.post(
            f"/api/forms/{form['id']}/submit", json={"file_ids": up}, headers=csrf(staff)
        )
        assert r.status_code == 409  # accepted rounds stay as they are

        # A company's own form: its file becomes a guideline its people may open.
        r = await client.post(
            "/api/forms",
            json={
                "name": "Rekod advance zon",
                "kind": "advance",
                "branch_id": b,
                "file_id": up[0],
                "schedule": {"every": "month", "from_day": 14, "to_day": 16},
            },
            headers=csrf(client),
        )
        assert r.status_code == 201, r.text
        assert r.json()["schedule"] == {"every": "month", "from_day": 14, "to_day": 16}
        r = await client.delete(
            f"/api/forms/{r.json()['id']}",
            headers={"content-type": "application/json", **csrf(client)},
        )
        assert r.status_code == 204
        assert len((await staff.get("/api/forms")).json()) == 1
    finally:
        await staff.aclose()
        await far.aclose()


async def test_ai_fills_the_form_into_the_persons_draft(client, llm, temporal):
    o = await office(client)
    b = o["branch"]["id"]
    staff = await as_role(
        client, "siti@example.com", "staff", branch_id=b, department_id=o["depts"]["Finance"]
    )
    try:
        twin = (await staff.post("/api/me/twin", json=ANSWERS, headers=csrf(staff))).json()["twin"]
        form = (
            await client.post(
                "/api/forms/starters",
                json={"key": "expense_claim", "branch_id": b},
                headers=csrf(client),
            )
        ).json()
        r = await staff.post(
            f"/api/forms/{form['id']}/ask",
            json={"text": "Grab ke PPD RM23.50 pada 3 Oktober; kertas A4 RM45."},
            headers=csrf(staff),
        )
        assert r.status_code == 201, r.text
        task_id = r.json()["task_id"]
        assert r.json()["agent"]["id"] == twin["id"]
        llm.call("describe_form", form_id=form["id"]).call(
            "fill_form",
            form_id=form["id"],
            fields={"Name": "Siti"},  # the form is built in the company's language
            rows=[
                {"Date": "2026-10-03", "Description": "Grab ke PPD", "Amount (RM)": "RM23.50"},
                {"Description": "Kertas A4", "Amount (RM)": 45},
            ],
        ).say("Diisi: 2 tuntutan, jumlah RM68.50.")
        step = await runtime.run_task_step(task_id)
        assert step.state == "done", step
        mine = (await staff.get("/api/forms")).json()[0]
        assert mine["state"] == "draft" and mine["mine"]["made_by"] == "agent"
        assert mine["mine"]["task_id"] == task_id and len(mine["mine"]["files"]) == 1
        fid = mine["mine"]["files"][0]["id"]
        filled = await staff.get(f"/api/files/{fid}/download")
        ws = load_workbook(io.BytesIO(filled.content)).worksheets[0]
        assert ws["C4"].value == "Siti" and ws["C8"].value == "Grab ke PPD"
        assert ws["E9"].value == 45
        # The person checks it and hands it in.
        r = await staff.post(
            f"/api/forms/{form['id']}/submit", json={"file_ids": [fid]}, headers=csrf(staff)
        )
        assert r.json()["state"] == "submitted"
        async with SessionLocal() as db:
            s = await db.get(FormSubmission, mine["mine"]["id"])
            assert s is not None and s.made_by == "agent" and s.status == "submitted"
    finally:
        await staff.aclose()
